import hashlib
import hmac
import io
import json
import os
import re
import secrets
import sqlite3
import threading
import time
import urllib.request
import urllib.parse
import uuid
from datetime import date
from decimal import Decimal, ROUND_HALF_UP, localcontext
from functools import wraps
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

import click
try:
    import psycopg
    PsycopgError = (psycopg.Error,)
except (ImportError, Exception):
    psycopg = None
    PsycopgError = ()
from flask import (
    Flask, abort, current_app, flash, g, jsonify, redirect, render_template,
    request, send_file, send_from_directory, session, url_for,
)
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from werkzeug.exceptions import HTTPException
from werkzeug.security import check_password_hash, generate_password_hash

from trade_parser import parse_text
from .database import get_db, install_database
from .prices import (
    DailyPriceScheduler, fetch_stock_quote, get_all_held_stocks,
    get_cached_prices, invalidate_prices_cache, update_all_stock_prices,
)
from .security import (
    configure_app, credential_version, csrf_token, local_setup_allowed, require_csrf,
    safe_redirect_target, secure_response,
)
from .storage import (
    DOWNLOAD_TTL, IMAGE_TYPES, MAX_CLOUD_IMAGE, MAX_LOCAL_IMAGE, UPLOAD_TTL,
    StorageUnavailable, get_storage, safe_path, validate_image,
)
from .validation import (
    COMMISSION_TYPES, MAX_ROWS, MAX_TEXT, ValidationError, client_name,
    identifier, number, rule_values, submission_uuid, text, trade_values, valid_date,
)


ZERO = Decimal("0")
CENT = Decimal("0.01")
TRADE_SELECT = """
SELECT t.*, c.name AS client_name, c.commission_type, c.commission_rate, u.filename AS upload_filename
FROM trades t JOIN clients c ON c.id=t.client_id
LEFT JOIN uploads u ON u.id=t.upload_id
"""
EXPORT_COLS = [
    ("Username", "client_name"), ("Stock", "stock"), ("Quantity", "quantity"),
    ("Buy Price", "buy_price"), ("Sell Price", "sell_price"), ("Buy Date", "buy_date"),
    ("Sell Date", "sell_date"), ("Status", "status"), ("Invested", "invested"),
    ("P&L", "pnl"), ("P&L %", "pnl_pct"), ("Commission", "commission"),
    ("Holding Days", "holding_days"), ("Source", "source"), ("Notes", "notes"),
]


_settings_cache = None
_settings_cache_time = 0.0
_settings_cache_ttl = 300.0  # 5 minutes
_settings_lock = threading.Lock()


def settings_values():
    global _settings_cache, _settings_cache_time
    if "settings_values" in g:
        return g.settings_values
    now = time.time()
    if _settings_cache is not None and (now - _settings_cache_time < _settings_cache_ttl):
        g.settings_values = dict(_settings_cache)
        return g.settings_values
    with _settings_lock:
        if _settings_cache is not None and (now - _settings_cache_time < _settings_cache_ttl):
            g.settings_values = dict(_settings_cache)
            return g.settings_values
        try:
            _settings_cache = {
                row["key"]: row["value"] for row in get_db().execute("SELECT key,value FROM settings").fetchall()
            }
            _settings_cache_time = now
        except Exception:
            _settings_cache = {}
        g.settings_values = dict(_settings_cache)
        return g.settings_values


def setting(key, default=None):
    return settings_values().get(key, default)


def set_setting(key, value):
    global _settings_cache, _settings_cache_time
    get_db().execute(
        "INSERT INTO settings(key,value) VALUES (?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (key, value),
    )
    with _settings_lock:
        if _settings_cache is not None:
            _settings_cache[key] = value
        _settings_cache_time = time.time()
    settings_values()[key] = value


def shorten_url(long_url):
    """Call the free is.gd API to get a short URL. Returns None on any failure."""
    try:
        api = "https://is.gd/create.php?format=simple&url=" + urllib.parse.quote(long_url, safe="")
        with urllib.request.urlopen(api, timeout=4) as resp:
            short = resp.read().decode().strip()
        if short.startswith("https://is.gd/") or short.startswith("http://is.gd/"):
            return short
    except Exception:
        pass
    return None


def portal_token():
    token = setting("upload_token")
    if not token:
        new_token = secrets.token_urlsafe(32)
        get_db().execute(
            "INSERT INTO settings(key,value) VALUES ('upload_token',?) ON CONFLICT(key) DO NOTHING",
            (new_token,),
        )
        g.pop("settings_values", None)
        token = setting("upload_token")
        # Generate a short URL (build directly to avoid circular portal_url call)
        try:
            base = current_app.config["PUBLIC_BASE_URL"]
            long = (base + url_for("public_upload", token=token)) if base else url_for("public_upload", token=token, _external=True)
            short = shorten_url(long)
            if short:
                set_setting("upload_short_url", short)
        except Exception:
            pass
    return token


def portal_url():
    token = portal_token()
    base = current_app.config["PUBLIC_BASE_URL"]
    if base:
        return base + url_for("public_upload", token=token)
    return url_for("public_upload", token=token, _external=True)


def admin_credentials():
    configured = current_app.config["ADMIN_PASSWORD_HASH"]
    if configured:
        return configured, current_app.config["ADMIN_AUTH_VERSION"]
    stored = setting("admin_password", "")
    return stored, credential_version(current_app.secret_key, "hash:" + stored) if stored else ""


def login_required(function):
    @wraps(function)
    def wrapper(*args, **kwargs):
        if not session.get("admin"):
            if session.get("client_id"):
                return redirect(url_for("client_portal"))
            if request.path.startswith("/api/"):
                abort(401, description="Administrator sign-in is required.")
            return redirect(url_for("login", next=request.path))
        return function(*args, **kwargs)
    return wrapper


def client_login_required(function):
    @wraps(function)
    def wrapper(*args, **kwargs):
        if not session.get("client_id"):
            if request.path.startswith("/api/"):
                abort(401, description="Client sign-in is required.")
            return redirect(url_for("login", next=request.path))
        return function(*args, **kwargs)
    return wrapper


def dec(value):
    return Decimal(str(value)) if value is not None else ZERO


def effective_rule(client):
    ctype = client["commission_type"] or setting("default_commission_type", "profit_pct")
    rate = client["commission_rate"]
    return ctype, dec(rate if rate is not None else setting("default_commission_rate", "10"))


_stock_aliases_ensured = False
_stock_aliases_cache = None  # in-memory dict: UPPER(alias) -> UPPER(canonical)
_stock_aliases_lock = threading.Lock()
_stock_suggestions_cache = None
_stock_suggestions_time = 0
_stock_suggestions_ttl = 300  # 5 minutes
_known_stocks_cache = None
_known_stocks_time = 0.0
_known_stocks_ttl = 300.0  # 5 minutes
_known_stocks_lock = threading.Lock()
_active_clients_cache = None
_active_clients_time = 0.0
_client_phones_cleaned = False


def ensure_stock_aliases_table(db):
    global _stock_aliases_ensured
    if _stock_aliases_ensured:
        return
    try:
        if getattr(db, "postgres", False):
            db.execute("""
                CREATE TABLE IF NOT EXISTS public.stock_aliases (
                    id BIGINT GENERATED BY DEFAULT AS IDENTITY PRIMARY KEY,
                    canonical_name TEXT NOT NULL,
                    alias TEXT NOT NULL UNIQUE,
                    created_at TIMESTAMPTZ NOT NULL DEFAULT timezone('utc'::text, now())
                )
            """)
            db.execute("CREATE INDEX IF NOT EXISTS idx_stock_aliases_canonical ON public.stock_aliases(canonical_name)")
            db.execute("CREATE INDEX IF NOT EXISTS idx_stock_aliases_alias ON public.stock_aliases(alias)")
        else:
            db.execute("""
                CREATE TABLE IF NOT EXISTS stock_aliases (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    canonical_name TEXT NOT NULL,
                    alias TEXT NOT NULL UNIQUE,
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                )
            """)
            db.execute("CREATE INDEX IF NOT EXISTS idx_stock_aliases_canonical ON stock_aliases(canonical_name)")
            db.execute("CREATE INDEX IF NOT EXISTS idx_stock_aliases_alias ON stock_aliases(alias)")
        _stock_aliases_ensured = True
    except Exception as exc:
        current_app.logger.warning("ensure_stock_aliases_table: %s", exc)


def invalidate_stock_caches():
    global _stock_aliases_cache, _stock_suggestions_cache, _stock_suggestions_time, _known_stocks_cache, _known_stocks_time
    with _stock_aliases_lock:
        _stock_aliases_cache = None
        _stock_suggestions_cache = None
        _stock_suggestions_time = 0
    with _known_stocks_lock:
        _known_stocks_cache = None
        _known_stocks_time = 0.0


def get_alias_mapping(db):
    global _stock_aliases_cache
    if _stock_aliases_cache is not None:
        return _stock_aliases_cache
    with _stock_aliases_lock:
        if _stock_aliases_cache is not None:
            return _stock_aliases_cache
        ensure_stock_aliases_table(db)
        mapping = {}
        try:
            rows = db.execute("SELECT canonical_name, alias FROM stock_aliases").fetchall()
            for r in rows:
                c = (r["canonical_name"] or "").strip().upper()
                a = (r["alias"] or "").strip().upper()
                if c:
                    mapping[c] = c
                if a and c:
                    mapping[a] = c
        except Exception:
            pass
        _stock_aliases_cache = mapping
        return _stock_aliases_cache


def resolve_canonical_stock(db, raw_stock):
    if not raw_stock:
        return raw_stock
    cleaned = raw_stock.strip().upper()
    mapping = get_alias_mapping(db)
    return mapping.get(cleaned, cleaned)


def get_known_stocks(db):
    global _known_stocks_cache, _known_stocks_time
    now = time.time()
    if _known_stocks_cache is not None and (now - _known_stocks_time < _known_stocks_ttl):
        return _known_stocks_cache

    with _known_stocks_lock:
        if _known_stocks_cache is not None and (now - _known_stocks_time < _known_stocks_ttl):
            return _known_stocks_cache
        ensure_stock_aliases_table(db)
        mapping = {}
        try:
            rows = db.execute("SELECT canonical_name, alias FROM stock_aliases ORDER BY canonical_name ASC, alias ASC").fetchall()
            for r in rows:
                c = (r["canonical_name"] or "").strip().upper()
                a = (r["alias"] or "").strip().upper()
                if c not in mapping:
                    mapping[c] = set()
                if a and a != c:
                    mapping[c].add(a)
        except Exception:
            pass

        result = []
        for c in sorted(mapping.keys()):
            alias_list = sorted(list(mapping[c]))
            result.append({
                "canonical_name": c,
                "aliases": ", ".join(alias_list),
                "alias_list": alias_list,
            })
        _known_stocks_cache = result
        _known_stocks_time = now
        return _known_stocks_cache


def get_client_entered_stocks(db):
    """Returns all distinct stock names found in trades, client holdings, and aliases with status metadata."""
    ensure_stock_aliases_table(db)
    all_stocks = set()
    try:
        rows = db.execute("SELECT DISTINCT stock FROM trades WHERE stock IS NOT NULL AND TRIM(stock) != ''").fetchall()
        for r in rows:
            s = r["stock"].strip().upper()
            if s:
                all_stocks.add(s)
    except Exception:
        pass

    try:
        rows = db.execute("SELECT DISTINCT stock FROM client_shares WHERE stock IS NOT NULL AND TRIM(stock) != ''").fetchall()
        for r in rows:
            s = r["stock"].strip().upper()
            if s:
                all_stocks.add(s)
    except Exception:
        pass

    try:
        rows = db.execute("SELECT DISTINCT canonical_name, alias FROM stock_aliases").fetchall()
        for r in rows:
            if r["canonical_name"]:
                all_stocks.add(r["canonical_name"].strip().upper())
            if r["alias"]:
                all_stocks.add(r["alias"].strip().upper())
    except Exception:
        pass

    # Map alias -> canonical_name and track canonical stocks
    alias_to_canonical = {}
    canonical_set = set()
    try:
        rows = db.execute("SELECT canonical_name, alias FROM stock_aliases").fetchall()
        for r in rows:
            c = r["canonical_name"].strip().upper()
            a = r["alias"].strip().upper()
            canonical_set.add(c)
            if a and a != c:
                alias_to_canonical[a] = c
    except Exception:
        pass

    stocks_list = []
    for s in sorted(all_stocks):
        mapped = alias_to_canonical.get(s)
        is_canonical = (s in canonical_set and mapped is None)
        stocks_list.append({
            "stock": s,
            "is_canonical": is_canonical,
            "mapped_to": mapped,
        })
    return stocks_list


POPULAR_STOCKS = [
    "ADANIENT", "ADANIPORTS", "ADANIGREEN", "ADANIPOWER", "APOLLOHOSP", "ASIANPAINT", "AXISBANK",
    "BAJAJ-AUTO", "BAJAJFINSV", "BAJFINANCE", "BEL", "BHEL", "BHARTIARTL", "BPCL",
    "CIPLA", "COALINDIA", "CHOLAFIN",
    "DABUR", "DIVISLAB", "DLF", "DRREDDY",
    "EICHERMOT",
    "FEDERALBNK",
    "GAIL", "GODREJCP", "GRASIM",
    "HAL", "HAVELLS", "HCLTECH", "HDFCBANK", "HDFCLIFE", "HEROMOTOCO", "HINDALCO", "HINDUNILVR",
    "ICICIBANK", "IDEA", "INDIGO", "INDUSINDBK", "INFY", "IOC", "IRFC", "ITC",
    "JINDALSTEL", "JIOFIN", "JSWSTEEL",
    "KOTAKBANK", "KPITTECH",
    "LT", "LTIM",
    "M&M", "MARUTI", "MOTHERSON",
    "NAUKRI", "NESTLEIND", "NHPC", "NTPC",
    "ONGC",
    "PAYTM", "PFC", "PIDILITIND", "PNB", "POWERGRID",
    "REC", "RELIANCE", "RVNL",
    "SAIL", "SBICARD", "SBILIFE", "SBIN", "SHREECEM", "SIEMENS", "SUNPHARMA", "SUZLON",
    "TATAMOTORS", "TATAPOWER", "TATASTEEL", "TCS", "TECHM", "TITAN", "TRENT",
    "ULTRACEMCO", "UPL",
    "VBL", "VEDL", "VOLTAS",
    "WIPRO",
    "YESBANK",
    "ZEEL", "ZOMATO"
]


def cleanup_stock_aliases_table(db):
    """Ensure that only main canonical stock names stay in stock_aliases, removing old variant aliases and updating ledger/holdings."""
    try:
        ensure_stock_aliases_table(db)
        rows = db.execute("SELECT DISTINCT canonical_name FROM stock_aliases WHERE canonical_name IS NOT NULL AND TRIM(canonical_name) != ''").fetchall()
        for r in rows:
            c = r["canonical_name"].strip().upper()
            if c:
                # Find all child aliases that are different from canonical
                child_aliases = db.execute("SELECT alias FROM stock_aliases WHERE canonical_name=? AND UPPER(TRIM(alias)) != ?", (c, c)).fetchall()
                for ca in child_aliases:
                    alias_name = ca["alias"].strip().upper()
                    if alias_name and alias_name != c:
                        db.execute("UPDATE trades SET stock=? WHERE UPPER(TRIM(stock))=?", (c, alias_name))
                        db.execute("UPDATE client_shares SET stock=? WHERE UPPER(TRIM(stock))=?", (c, alias_name))
                        db.execute("UPDATE client_alerts SET stock=? WHERE UPPER(TRIM(stock))=?", (c, alias_name))
                        try:
                            db.execute("DELETE FROM stock_prices WHERE UPPER(TRIM(stock))=?", (alias_name,))
                        except Exception:
                            pass
                # Delete all child alias rows so only the main name stays
                db.execute("DELETE FROM stock_aliases WHERE canonical_name=? AND UPPER(TRIM(alias)) != ?", (c, c))
                # Ensure main canonical name exists as (canonical, canonical)
                db.execute(
                    """INSERT INTO stock_aliases (canonical_name, alias) VALUES (?, ?)
                       ON CONFLICT(alias) DO UPDATE SET canonical_name=excluded.canonical_name""",
                    (c, c)
                )
        invalidate_stock_caches()
    except Exception as exc:
        current_app.logger.warning("cleanup_stock_aliases_table error: %s", exc)


def get_stock_suggestions(db):
    """Returns a list of deduplicated main stock suggestion dicts (cached in memory for speed)."""
    global _stock_suggestions_cache, _stock_suggestions_time
    now = time.time()
    if _stock_suggestions_cache is not None and (now - _stock_suggestions_time < _stock_suggestions_ttl):
        return _stock_suggestions_cache

    alias_map = get_alias_mapping(db)
    items_map = {}

    # 1. Configured canonical stocks from stock_aliases
    try:
        rows = db.execute("SELECT DISTINCT canonical_name FROM stock_aliases WHERE canonical_name IS NOT NULL AND TRIM(canonical_name) != '' ORDER BY canonical_name ASC").fetchall()
        for r in rows:
            c = (r["canonical_name"] or "").strip().upper()
            if c and c not in items_map:
                items_map[c] = {"symbol": c, "canonical": c}
    except Exception:
        pass

    # 2. Distinct stocks from trades, client_shares, and stock_prices in a single UNION query
    try:
        rows = db.execute(
            """
            SELECT DISTINCT stock FROM (
                SELECT stock FROM trades WHERE stock IS NOT NULL AND TRIM(stock) != ''
                UNION
                SELECT stock FROM client_shares WHERE stock IS NOT NULL AND TRIM(stock) != ''
                UNION
                SELECT stock FROM stock_prices WHERE stock IS NOT NULL AND TRIM(stock) != ''
            ) u
            """
        ).fetchall()
        for r in rows:
            raw_s = (r["stock"] or "").strip().upper()
            canon = alias_map.get(raw_s, raw_s)
            if canon and canon not in items_map:
                items_map[canon] = {"symbol": canon, "canonical": canon}
    except Exception:
        pass

    # 3. Common market stocks (resolved to canonical name in memory)
    for sym in POPULAR_STOCKS:
        canon = alias_map.get(sym, sym)
        if canon not in items_map:
            items_map[canon] = {"symbol": canon, "canonical": canon}

    res = sorted(items_map.values(), key=lambda x: x["symbol"])
    _stock_suggestions_cache = res
    _stock_suggestions_time = now
    return res


def clean_phone_number(raw_phone):
    """Strip any +91, 91, or 0 country codes and keep only the 10-digit mobile number."""
    if not raw_phone:
        return ""
    s = str(raw_phone).strip()
    if s.startswith("+91"):
        s = s[3:]
    elif s.startswith("+"):
        s = s[1:]
    digits = re.sub(r"\D", "", s)
    if digits.startswith("91") and len(digits) > 10:
        digits = digits[2:]
    while digits.startswith("0") and len(digits) > 10:
        digits = digits[1:]
    if len(digits) > 10:
        digits = digits[-10:]
    return digits


def cleanup_existing_client_phones(db):
    """Ensure any client phones with +91 prefixes in SQLite/Supabase are normalized to 10 digits."""
    try:
        rows = db.execute("SELECT id, phone FROM clients WHERE phone IS NOT NULL AND phone != ''").fetchall()
        for r in rows:
            clean = clean_phone_number(r["phone"])
            if clean and clean != r["phone"]:
                conflict = db.execute("SELECT id FROM clients WHERE phone=? AND id!=?", (clean, r["id"])).fetchone()
                if not conflict:
                    db.execute("UPDATE clients SET phone=? WHERE id=?", (clean, r["id"]))
    except Exception:
        pass


def get_client(cid):
    client = get_db().execute("SELECT * FROM clients WHERE id=?", (cid,)).fetchone()
    if not client:
        abort(404, description="That user no longer exists.")
    return client


def get_or_create_client(name, phone=None):
    name = client_name(name)
    database = get_db()
    if phone:
        phone = clean_phone_number(phone)
        if phone:
            row = database.execute("SELECT id FROM clients WHERE phone=?", (phone,)).fetchone()
            if row:
                return row["id"]
            row = database.execute(
                "INSERT INTO clients(name, phone) VALUES (?, ?) ON CONFLICT DO NOTHING RETURNING id", (name, phone)
            ).fetchone()
            if row:
                return row["id"]

    row = database.execute("SELECT id FROM clients WHERE lower(name)=lower(?)", (name,)).fetchone()
    if row:
        return row["id"]
    row = database.execute(
        "INSERT INTO clients(name) VALUES (?) ON CONFLICT DO NOTHING RETURNING id", (name,)
    ).fetchone()
    if row:
        return row["id"]
    return database.execute("SELECT id FROM clients WHERE lower(name)=lower(?)", (name,)).fetchone()["id"]


def insert_trade(cid, trade, source="manual", upload_id=None):
    return insert_trades(cid, [trade], source, upload_id)[0]


def insert_trades(cid, trades, source="manual", upload_id=None):
    if not trades or len(trades) > MAX_ROWS:
        raise ValidationError(f"Save between 1 and {MAX_ROWS} trades in a batch.")
    if source not in ("manual", "import", "upload"):
        source = "manual"
    ctype, rate = effective_rule(get_client(cid))
    parameters = []
    db = get_db()
    for trade in trades:
        stock = resolve_canonical_stock(db, trade["stock"])
        bp = trade.get("buy_price")
        sp = trade.get("sell_price")
        b_date = trade.get("buy_date")
        s_date = trade.get("sell_date")

        # Defensive pairs check to strictly satisfy DB constraints:
        # trades_buy_pair: ((buy_price IS NULL) = (buy_date IS NULL))
        if bp is not None:
            if not b_date:
                b_date = date.today().isoformat()
        else:
            b_date = None

        # trades_sell_pair: ((sell_price IS NULL) = (sell_date IS NULL))
        if sp is not None:
            if not s_date:
                s_date = date.today().isoformat()
        else:
            s_date = None

        # trades_date_order: sell_date >= buy_date if both are present
        if b_date and s_date and s_date < b_date:
            s_date = b_date

        parameters.extend((
            cid, stock, trade["quantity"], bp, sp,
            b_date, s_date, trade.get("commission_override"),
            ctype, rate, source, upload_id, trade.get("notes"),
        ))
    rows = db.execute(
        """INSERT INTO trades(
            client_id,stock,quantity,buy_price,sell_price,buy_date,sell_date,commission_override,
            commission_type_snapshot,commission_rate_snapshot,source,upload_id,notes
        ) VALUES """ + ",".join("(?,?,?,?,?,?,?,?,?,?,?,?,?)" for _ in trades) + " RETURNING id",
        parameters,
    ).fetchall()
    return [row["id"] for row in rows]


def enrich(trade):
    result = dict(trade)
    for key in ("quantity", "buy_price", "sell_price", "commission_rate",
                "commission_override", "commission_rate_snapshot"):
        if result.get(key) is not None:
            result[key] = dec(result[key])
    quantity = result["quantity"]
    closed = result["buy_price"] is not None and result["sell_price"] is not None
    result["status"] = "Closed" if closed else "Open"
    result["invested"] = (result["buy_price"] or ZERO) * quantity
    result["pnl"] = (result["sell_price"] - result["buy_price"]) * quantity if closed else ZERO
    result["pnl_pct"] = result["pnl"] / result["invested"] * 100 if closed and result["invested"] else ZERO
    if result.get("commission_type_snapshot") is not None:
        ctype, rate = result["commission_type_snapshot"], result["commission_rate_snapshot"]
    else:
        ctype, rate = effective_rule(result)
    result["eff_type"], result["eff_rate"] = ctype, rate
    if result.get("commission_override") is not None:
        commission = result["commission_override"]
    elif not closed:
        commission = ZERO
    elif ctype == "turnover_pct":
        commission = (result["buy_price"] + result["sell_price"]) * quantity * rate / 100
    elif ctype == "flat":
        commission = rate
    else:
        commission = max(result["pnl"], ZERO) * rate / 100
    result["commission"] = commission.quantize(CENT, rounding=ROUND_HALF_UP)
    result["holding_days"] = None
    if result["buy_date"] and result["sell_date"]:
        try:
            result["holding_days"] = (
                date.fromisoformat(result["sell_date"]) - date.fromisoformat(result["buy_date"])
            ).days
        except ValueError:
            current_app.logger.warning("Legacy trade %s has an invalid stored date.", result["id"])
    return result


def query_trades(args):
    where, parameters = [], []
    if args.get("client_id"):
        where.append("t.client_id=?")
        parameters.append(identifier(args["client_id"]))
    if args.get("stock"):
        where.append("lower(t.stock) LIKE lower(?)")
        stock = text(args["stock"], "Stock filter", limit=120)
        parameters.append("%" + stock + "%")
    start = valid_date(args.get("from"), "From date")
    end = valid_date(args.get("to"), "To date")
    if start and end and start > end:
        raise ValidationError("From date cannot be after to date.")
    if start:
        where.append("COALESCE(t.sell_date,t.buy_date)>=?")
        parameters.append(start)
    if end:
        where.append("COALESCE(t.sell_date,t.buy_date)<=?")
        parameters.append(end)
    status = args.get("status", "")
    if status not in ("", None, "open", "closed"):
        raise ValidationError("Choose a valid trade status.")
    if status == "open":
        where.append("(t.buy_price IS NULL OR t.sell_price IS NULL)")
    elif status == "closed":
        where.append("t.buy_price IS NOT NULL AND t.sell_price IS NOT NULL")
    sql = TRADE_SELECT + (" WHERE " + " AND ".join(where) if where else "")
    sql += " ORDER BY COALESCE(t.sell_date,t.buy_date) DESC NULLS LAST,t.id DESC"
    return [enrich(row) for row in get_db().execute(sql, parameters).fetchall()]


def summarize(trades):
    closed = [trade for trade in trades if trade["status"] == "Closed"]
    return {
        "count": len(trades), "closed": len(closed), "open": len(trades) - len(closed),
        "pnl": sum((trade["pnl"] for trade in closed), ZERO),
        "commission": sum((trade["commission"] for trade in trades), ZERO),
        "invested_open": sum((trade["invested"] for trade in trades if trade["status"] == "Open"), ZERO),
        "wins": sum(trade["pnl"] > 0 for trade in closed),
        "losses": sum(trade["pnl"] < 0 for trade in closed),
        "turnover": sum(
            (((trade["buy_price"] or ZERO) + (trade["sell_price"] or ZERO)) * trade["quantity"]
             for trade in trades), ZERO
        ),
    }


def payments_by_client():
    rows = get_db().execute("SELECT client_id,SUM(amount) AS total FROM payments GROUP BY client_id").fetchall()
    return {row["client_id"]: dec(row["total"]) for row in rows}


def funds_by_client():
    rows = get_db().execute("SELECT client_id,SUM(amount) AS total FROM funds GROUP BY client_id").fetchall()
    return {row["client_id"]: dec(row["total"]) for row in rows}


def sync_client_holdings_from_trades(cid=None, db=None):
    """Sync open buy trades into client_shares in batch so client_shares remains the single canonical store."""
    if db is None:
        db = get_db()

    where_sql = "WHERE client_id=?" if cid is not None else ""
    params = (cid,) if cid is not None else ()

    try:
        open_trades = db.execute(
            f"""SELECT id, client_id, stock, quantity, buy_price, COALESCE(buy_date, created_at) as added_on
                FROM trades
                {where_sql} {"AND" if where_sql else "WHERE"} buy_price IS NOT NULL AND (sell_price IS NULL OR sell_price=0)
                ORDER BY id ASC""",
            params
        ).fetchall()
        if not open_trades:
            return

        client_ids = list({ot["client_id"] for ot in open_trades})
        placeholders = ",".join("?" for _ in client_ids)
        existing_rows = db.execute(
            f"SELECT client_id, UPPER(TRIM(stock)) as stock FROM client_shares WHERE client_id IN ({placeholders})",
            client_ids
        ).fetchall()
        existing_set = {(r["client_id"], r["stock"]) for r in existing_rows}
        alias_map = get_alias_mapping(db)

        to_insert = []
        for ot in open_trades:
            c_id = ot["client_id"]
            raw_s = (ot["stock"] or "").strip().upper()
            stk = alias_map.get(raw_s, raw_s)
            if not stk:
                continue
            if (c_id, stk) not in existing_set:
                qty = int(round(dec(ot["quantity"])))
                bp = dec(ot["buy_price"]).quantize(CENT, rounding=ROUND_HALF_UP)
                added_on = ot["added_on"] or date.today().isoformat()
                if isinstance(added_on, (date, datetime)):
                    added_on = added_on.isoformat()
                elif len(str(added_on)) > 10:
                    added_on = str(added_on)[:10]
                to_insert.append((c_id, stk, qty, bp, added_on))
                existing_set.add((c_id, stk))

        for item in to_insert:
            db.execute(
                "INSERT INTO client_shares (client_id, stock, quantity, buy_price, added_on) VALUES (?, ?, ?, ?, ?)",
                item
            )
        if to_insert:
            db.commit()
    except Exception as exc:
        current_app.logger.warning("sync_client_holdings_from_trades: %s", exc)


def get_client_calculated_holdings(cid, db=None, sync_trades=False):
    if db is None:
        db = get_db()

    if sync_trades:
        sync_client_holdings_from_trades(cid, db)

    try:
        raw_shares = db.execute(
            "SELECT id, stock, quantity, buy_price, added_on FROM client_shares WHERE client_id=? ORDER BY stock ASC",
            (cid,)
        ).fetchall()
    except Exception:
        raw_shares = []

    alias_map = get_alias_mapping(db)
    results = []
    for r in raw_shares:
        raw_s = (r["stock"] or "").strip().upper()
        stk = alias_map.get(raw_s, raw_s)
        if not stk:
            continue
        qty = dec(r["quantity"]) if r["quantity"] is not None else ZERO
        bp = dec(r["buy_price"]) if r["buy_price"] is not None else ZERO
        if qty <= ZERO:
            continue
        tot_invested = (qty * bp).quantize(CENT, rounding=ROUND_HALF_UP)
        results.append({
            "id": r["id"],
            "stock": stk,
            "quantity": int(round(qty)),
            "buy_price": bp.quantize(CENT, rounding=ROUND_HALF_UP),
            "total_invested": tot_invested,
            "added_on": str(r["added_on"]) if r.get("added_on") else "",
        })

    results.sort(key=lambda x: x["stock"])
    return results


def get_clients_holding_stock(stock, db=None):
    """Returns list of active clients ({'id': ..., 'name': ...}) who currently hold the specified stock (canonical-aware)."""
    if db is None:
        db = get_db()
    if not stock:
        return []
    target_canonical = resolve_canonical_stock(db, stock)
    alias_map = get_alias_mapping(db)
    try:
        rows = db.execute(
            """
            SELECT DISTINCT c.id, c.name, h.stock, h.quantity
            FROM client_shares h
            JOIN clients c ON c.id = h.client_id
            WHERE c.active = 1 AND h.quantity > 0
            ORDER BY lower(c.name)
            """
        ).fetchall()
        holding_clients = []
        seen = set()
        for r in rows:
            sym = (r["stock"] or "").strip().upper()
            if alias_map.get(sym, sym) == target_canonical and r["id"] not in seen:
                seen.add(r["id"])
                holding_clients.append({"id": r["id"], "name": r["name"]})
        return holding_clients
    except Exception:
        return []


def sync_trade_profits_to_funds(db, cid=None):
    """Ensure completed trades have their realized profit recorded in funds ledger exactly once."""
    try:
        where_clause = "WHERE buy_price IS NOT NULL AND sell_price IS NOT NULL"
        params = []
        if cid:
            where_clause += " AND client_id=?"
            params.append(cid)
        closed_trades = db.execute(
            f"SELECT id, client_id, stock, quantity, buy_price, sell_price, sell_date, created_at FROM trades {where_clause}",
            params
        ).fetchall()
        if not closed_trades:
            return

        funds_params = ["%[Trade #%"] + ([cid] if cid else [])
        funds_sql = "SELECT notes FROM funds WHERE notes LIKE ?" + (" AND client_id=?" if cid else "")
        existing_notes_rows = db.execute(funds_sql, funds_params).fetchall()
        existing_notes_text = " ".join([r["notes"] or "" for r in existing_notes_rows])

        for t in closed_trades:
            pnl = ((dec(t["sell_price"]) - dec(t["buy_price"])) * dec(t["quantity"])).quantize(CENT, rounding=ROUND_HALF_UP)
            if pnl == ZERO:
                continue
            tag = f"[Trade #{t['id']}]"
            if tag not in existing_notes_text:
                note = f"[Realized {'Profit' if pnl > 0 else 'Loss'}] {t['stock']} ({t['quantity']} sh) {tag}"
                date_val = t["sell_date"] or str(t["created_at"])[:10] or date.today().isoformat()
                db.execute(
                    "INSERT INTO funds (client_id, amount, added_on, notes) VALUES (?, ?, ?, ?)",
                    (t["client_id"], pnl, date_val, note)
                )
                existing_notes_text += " " + note
    except Exception as exc:
        current_app.logger.warning("sync_trade_profits_to_funds error: %s", exc)



def rule_label(ctype, rate):
    if not ctype or rate is None:
        return "Default"
    r = f"{dec(rate):.1f}".rstrip('0').rstrip('.')
    if ctype == "flat":
        return f"{r} flat/trade"
    return f"{r}{COMMISSION_TYPES[ctype]}"


def client_values(values, name):
    ctype, rate = rule_values(values.get("commission_type"), values.get("commission_rate"), allow_default=True)
    raw_phone = values.get("phone")
    phone = clean_phone_number(raw_phone)
    if not phone or len(phone) != 10:
        raise ValidationError("Please provide a valid 10-digit mobile number.")
    return (
        name, phone,
        text(values.get("email"), "Email", limit=254) or None, ctype, rate,
        text(values.get("notes"), "Notes", multiline=True) or None,
        number(values.get("portfolio_amount"), "Portfolio amount", positive=True) if values.get("portfolio_amount") else ZERO,
    )


def upload_records(cid=None):
    sql = "SELECT u.*,c.name AS current_client_name FROM uploads u LEFT JOIN clients c ON c.id=u.client_id"
    parameters = ()
    if cid is not None:
        sql += " WHERE u.client_id=?"
        parameters = (cid,)
    rows = get_db().execute(sql + " ORDER BY u.id DESC", parameters).fetchall()
    for row in rows:
        row["entered_name"] = row["client_name"]
        row["client_name"] = row.pop("current_client_name") or row["client_name"]
    return rows


def sheet(worksheet, headers, rows):
    worksheet.append(headers)
    for row in rows:
        worksheet.append([
            ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value for value in row
        ])
        for cell, value in zip(worksheet[worksheet.max_row], row):
            if isinstance(value, str):
                cell.data_type = "s"
                cell.quotePrefix = True
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for column in worksheet.columns:
        width = max(len(str(cell.value or "")) for cell in column) + 2
        worksheet.column_dimensions[column[0].column_letter].width = min(max(width, 10), 40)
    worksheet.freeze_panes = "A2"


def workbook_response(workbook, filename):
    buffer = io.BytesIO()
    workbook.save(buffer)
    if current_app.config["PRODUCTION"] and buffer.tell() > 4 * 1024 * 1024:
        abort(413, description="This export exceeds the cloud response limit. Export a smaller user or date range.")
    buffer.seek(0)
    return send_file(
        buffer, as_attachment=True, download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def json_object():
    if not request.is_json:
        abort(415, description="Send an application/json request.")
    data = request.get_json()
    if not isinstance(data, dict):
        raise ValidationError("The request body must be a JSON object.")
    return data


def owner_hash():
    if "_upload_owner" not in session:
        session["_upload_owner"] = secrets.token_urlsafe(32)
    return hashlib.sha256(session["_upload_owner"].encode()).hexdigest()


def scope_hash(mode):
    if mode == "admin":
        if not session.get("admin"):
            abort(403, description="Administrator authorization has expired.")
        source = "admin:" + admin_credentials()[1]
    elif mode == "portal":
        token = setting("upload_token", "")
        if not token:
            abort(403, description="This submission link is inactive.")
        source = "portal:" + token + ":" + setting("upload_code", "")
    else:
        abort(403, description="Invalid upload authorization.")
    return credential_version(current_app.secret_key, "upload-scope:" + source)


def authorize_upload(data):
    token = text(data.get("portal_token"), "Portal token", limit=128)
    code = text(data.get("code"), "Upload code", limit=128)
    if token:
        expected = setting("upload_token", "")
        if not expected or not hmac.compare_digest(token.encode(), expected.encode()):
            abort(403, description="This submission link is invalid or has been replaced.")
        expected_code = setting("upload_code", "")
        if expected_code and not hmac.compare_digest(code.encode(), expected_code.encode()):
            abort(403, description="Invalid upload code.")
        mode = "portal"
    elif session.get("admin"):
        mode = "admin"
    else:
        abort(403, description="Open your private submission link to upload trades.")
    return {"owner": owner_hash(), "scope": scope_hash(mode), "mode": mode}


def receipt_signer():
    return URLSafeTimedSerializer(current_app.secret_key, salt="commission-upload-v1")


def payment_signer():
    return URLSafeTimedSerializer(current_app.secret_key, salt="commission-pay-v1")


def decode_receipt(receipt):
    if not isinstance(receipt, str) or len(receipt) > 2048:
        raise ValidationError("A signed receipt is required for every image.")
    try:
        data = receipt_signer().loads(receipt, max_age=UPLOAD_TTL)
    except SignatureExpired as exc:
        raise ValidationError("The image upload authorization expired. Upload the image again.") from exc
    except BadSignature as exc:
        raise ValidationError("Invalid image upload receipt.") from exc
    if not isinstance(data, dict):
        raise ValidationError("Invalid image upload receipt.")
    return data


def receipt_asset(path, receipt, authorization=None):
    path = safe_path(path)
    decoded = decode_receipt(receipt)
    if authorization is None:
        authorization = {
            "owner": owner_hash(), "scope": scope_hash(decoded.get("mode")), "mode": decoded.get("mode"),
        }
    if (decoded.get("path") != path or decoded.get("owner") != authorization["owner"]
            or decoded.get("scope") != authorization["scope"]):
        abort(403, description="This image belongs to a different upload session or an expired link.")
    asset = get_db().execute("SELECT * FROM upload_assets WHERE path=?", (path,)).fetchone()
    if (not asset or asset["owner_hash"] != authorization["owner"]
            or asset["scope_hash"] != authorization["scope"] or asset["expires_at"] < int(time.time())):
        abort(403, description="This image upload authorization is no longer valid.")
    if asset["upload_id"] is not None:
        abort(409, description="This image has already been attached to a submission.")
    return asset


def prior_submission(submission_id, digest, owner):
    row = get_db().execute(
        "SELECT trade_count,submission_hash,submission_owner FROM uploads WHERE submission_id=?",
        (submission_id,),
    ).fetchone()
    if row:
        if row["submission_owner"] != owner or row["submission_hash"] != digest:
            return "CONFLICT"
        return jsonify(ok=True, saved=row["trade_count"], submission_id=submission_id)
    return None


def create_app(config=None):
    if (not (config and config.get("TESTING")) and os.environ.get("VERCEL") != "1"
            and os.environ.get("COMMISSION_ENV") != "production"):
        from dotenv import load_dotenv
        load_dotenv(Path(__file__).resolve().parents[1] / ".env", override=False)
    root = str(Path(__file__).resolve().parents[1])
    app = Flask(__name__, root_path=root, template_folder="templates", static_folder="static")
    configure_app(app, config)
    install_database(app)
    app.after_request(secure_response)
    app.jinja_env.globals["csrf_token"] = csrf_token

    @app.context_processor
    def inject_admin_context():
        if session.get("admin"):
            global _active_clients_cache, _active_clients_time
            now = time.time()
            if _active_clients_cache is not None and (now - _active_clients_time < 60.0):
                return {"active_clients": _active_clients_cache}
            try:
                clients = get_db().execute("SELECT id, name FROM clients WHERE active=1 ORDER BY lower(name)").fetchall()
                _active_clients_cache = clients
                _active_clients_time = now
                return {"active_clients": clients}
            except Exception:
                return {}
        return {}

    @app.before_request
    def authorize_request():
        g.decimal_context = localcontext()
        g.decimal_context.__enter__().prec = 60
        if session.get("client_id") or session.get("admin"):
            session.permanent = True
        if session.get("admin"):
            _, version = admin_credentials()
            if not version or session.get("auth_version") != version:
                session.clear()
        require_csrf()

    @app.teardown_request
    def restore_decimal_context(_error=None):
        context = g.pop("decimal_context", None)
        if context is not None:
            context.__exit__(None, None, None)

    @app.template_filter("money")
    def money(value):
        if value is None or value == "":
            return "\u2014"
        value = dec(value)
        symbol = setting("currency", "\u20b9")
        sign = "-" if value < 0 else ""
        whole, fraction = f"{abs(value):.2f}".split(".")
        if len(whole) > 3:
            head, tail = whole[:-3], whole[-3:]
            groups = []
            while len(head) > 2:
                groups.insert(0, head[-2:])
                head = head[:-2]
            whole = ",".join(([head] if head else []) + groups + [tail])
        return f"{sign}{symbol}{whole}.{fraction}"

    @app.template_filter("num")
    def num(value):
        if value is None or value == "":
            return ""
        value = dec(value)
        return f"{value:,.0f}" if value == value.to_integral_value() else f"{value:,.6f}".rstrip("0").rstrip(".")

    @app.context_processor
    def inject_globals():
        error = g.get("error_response", False)
        return {
            "COMMISSION_TYPES": COMMISSION_TYPES, "logged_in": bool(session.get("admin")),
            "is_admin": bool(session.get("admin")),
            "is_client": bool(session.get("client_id")),
            "client_name": session.get("client_name"),
            "client_id": session.get("client_id"),
            "today": date.today().isoformat(), "app_name": "Commission Book",
            "business_name": "" if error else setting("business_name", ""),
            "currency": "\u20b9" if error else setting("currency", "\u20b9"),
            "backend_mode": app.config["BACKEND_MODE"],
            "deployed": app.config["DEPLOYED"],
            "password_managed": bool(app.config["ADMIN_PASSWORD_HASH"]),
            "admin_password_managed": bool(app.config["ADMIN_PASSWORD_HASH"]),
            "commission_policy": "snapshot",
            "upload_url": portal_url() if session.get("admin") and not error else None,
            "all_stock_suggestions": [],
            "known_stocks": [],
        }

    def error_response(status, title, message):
        if request.path.startswith("/api/"):
            return jsonify(error=message), status
        g.error_response = True
        return render_template("error.html", status=status, title=title, message=message), status

    @app.errorhandler(ValidationError)
    def invalid_input(error):
        return error_response(400, "Please check your entry", str(error))

    @app.errorhandler(HTTPException)
    def http_error(error):
        message = error.description
        if error.code == 413 and request.method != "GET":
            message = "The request is too large. Upload images separately (4 MB locally; 10 MB directly to cloud storage)."
        return error_response(error.code, error.name, message)

    @app.errorhandler(StorageUnavailable)
    def unavailable_storage(error):
        return error_response(503, "Private storage unavailable", str(error))

    @app.errorhandler(sqlite3.IntegrityError)
    def database_conflict(error):
        app.logger.warning("A database constraint rejected a request (%s).", type(error).__name__)
        return error_response(409, "Conflicting entry", "This entry conflicts with existing data. Reload and try again.")

    if psycopg is not None:
        app.errorhandler(psycopg.IntegrityError)(database_conflict)

    @app.errorhandler(sqlite3.OperationalError)
    def unavailable_database(error):
        app.logger.error("Database operation failed (%s).", type(error).__name__)
        return error_response(503, "Ledger unavailable", "The database is unavailable. Try again shortly or check its configuration.")

    if psycopg is not None:
        app.errorhandler(psycopg.OperationalError)(unavailable_database)

    @app.errorhandler(500)
    def server_error(_error):
        return error_response(500, "Something went wrong", "Your request could not be completed. Please try again.")

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if session.get("admin"):
            return redirect(url_for("dashboard"))
        if session.get("client_id"):
            return redirect(url_for("client_portal"))
        password_hash, version = admin_credentials()
        first_run = not password_hash
        if first_run and not local_setup_allowed():
            abort(503, description="An administrator must configure the password on the server before this site can be used.")
        if request.method == "POST":
            login_type = request.form.get("login_type", "admin")
            if login_type == "client":
                identifier = request.form.get("identifier", "").strip()
                clean_ident = identifier.lstrip("@").strip()
                clean_phone = clean_phone_number(identifier)
                password = request.form.get("client_password", "")
                if not identifier or not password:
                    flash("Please provide your ID/phone and password.", "error")
                    return render_template("login.html", first_run=first_run, active_tab="client", public_token=portal_token()), 400
                client = get_db().execute(
                    """SELECT * FROM clients 
                    WHERE (phone=? OR phone=? OR lower(phone)=lower(?) OR lower(email)=lower(?) OR lower(username)=lower(?) OR lower(username)=lower(?) OR lower(name)=lower(?) OR CAST(id AS TEXT)=?) 
                      AND active=1""",
                    (identifier, clean_phone, identifier, identifier, identifier, clean_ident, identifier, identifier),
                ).fetchone()
                if not client or not client.get("password_hash"):
                    flash("Invalid credentials or client password has not been configured yet.", "error")
                    return render_template("login.html", first_run=first_run, active_tab="client", public_token=portal_token()), 401
                if not check_password_hash(client["password_hash"], password):
                    flash("Incorrect password.", "error")
                    return render_template("login.html", first_run=first_run, active_tab="client", public_token=portal_token()), 401
                session.clear()
                session["client_id"] = client["id"]
                session["client_name"] = client["name"]
                session["role"] = "client"
                session.permanent = True
                return redirect(safe_redirect_target(request.args.get("next"), url_for("client_portal")))
            else:
                password = request.form.get("password", "")
                if len(password) > 1024:
                    raise ValidationError("Password is too long.")
                if first_run:
                    if len(password) < 10 or password != request.form.get("confirm"):
                        flash("Password must be at least 10 characters and match confirmation.", "error")
                        return render_template("login.html", first_run=True, active_tab="admin"), 400
                    password_hash = generate_password_hash(password)
                    set_setting("admin_password", password_hash)
                    version = admin_credentials()[1]
                elif not check_password_hash(password_hash, password):
                    flash("Incorrect password.", "error")
                    return render_template("login.html", first_run=False, active_tab="admin"), 401
                session.clear()
                session["admin"], session["auth_version"] = True, version
                session.permanent = True
                return redirect(safe_redirect_target(request.args.get("next"), url_for("dashboard")))
        return render_template("login.html", first_run=first_run, active_tab=request.args.get("tab", "options"), public_token=portal_token())

    @app.route("/logout", methods=["POST"])
    def logout():
        session.clear()
        return redirect(url_for("login"))

    @app.route("/")
    @login_required
    def dashboard():
        args = {key: request.args.get(key, "") for key in ("from", "to")}
        trades = query_trades(args)
        total, paid = summarize(trades), payments_by_client()
        clients = get_db().execute("SELECT * FROM clients ORDER BY lower(name),id").fetchall()
        all_trades = query_trades({}) if any(args.values()) else trades
        all_earned = {}
        for trade in all_trades:
            all_earned[trade["client_id"]] = all_earned.get(trade["client_id"], ZERO) + trade["commission"]
        per_client = []
        for client in clients:
            summary = summarize([trade for trade in trades if trade["client_id"] == client["id"]])
            summary.update(
                id=client["id"], name=client["name"], paid=paid.get(client["id"], ZERO),
                outstanding=all_earned.get(client["id"], ZERO) - paid.get(client["id"], ZERO),
                rule=rule_label(client["commission_type"], client["commission_rate"]),
            )
            per_client.append(summary)
        per_client.sort(key=lambda value: value["commission"], reverse=True)
        total["paid"] = sum(paid.values(), ZERO)
        total["outstanding"] = sum((client["outstanding"] for client in per_client), ZERO)
        monthly, stocks = {}, {}
        for trade in trades:
            if trade["status"] == "Closed":
                month = (trade["sell_date"] or trade["buy_date"] or "")[:7] or "Unknown"
                entry = monthly.setdefault(month, {"commission": ZERO, "pnl": ZERO})
                entry["commission"] += trade["commission"]
                entry["pnl"] += trade["pnl"]
            stock = trade["stock"].upper()
            entry = stocks.setdefault(stock, {"stock": stock, "count": 0, "pnl": ZERO, "commission": ZERO})
            entry["count"] += 1
            entry["pnl"] += trade["pnl"]
            entry["commission"] += trade["commission"]
        months = sorted(monthly)
        positive = [client for client in per_client if client["commission"] > 0]
        chart = {
            "months": months, "currency": setting("currency", "\u20b9"),
            "commission": [float(round(monthly[month]["commission"], 2)) for month in months],
            "pnl": [float(round(monthly[month]["pnl"], 2)) for month in months],
            "clients": [client["name"] for client in positive],
            "client_commission": [float(round(client["commission"], 2)) for client in positive],
        }
        return render_template(
            "dashboard.html", total=total, per_client=per_client, chart=chart,
            top_stocks=sorted(stocks.values(), key=lambda stock: stock["commission"], reverse=True)[:10],
            recent=trades[:8], args=args,
        )

    @app.route("/trades")
    @login_required
    def trades_list():
        args = {key: request.args.get(key, "") for key in ("client_id", "stock", "from", "to", "status")}
        trades = query_trades(args)
        clients = get_db().execute("SELECT id,name FROM clients ORDER BY lower(name)").fetchall()
        return render_template("trades.html", trades=trades, total=summarize(trades), clients=clients, args=args)

    @app.route("/trades/new", methods=["GET", "POST"])
    @app.route("/trades/<int:tid>/edit", methods=["GET", "POST"])
    @login_required
    def trade_form(tid=None):
        trade = get_db().execute(TRADE_SELECT + " WHERE t.id=?", (tid,)).fetchone() if tid else None
        if tid and not trade:
            abort(404)
        if request.method == "POST":
            values = trade_values(request.form)
            name = client_name(request.form.get("client_name"))
            cid = get_or_create_client(name)
            if tid:
                get_db().execute(
                    """UPDATE trades SET client_id=?,stock=?,quantity=?,buy_price=?,sell_price=?,
                    buy_date=?,sell_date=?,commission_override=?,notes=? WHERE id=?""",
                    (cid, values["stock"], values["quantity"], values["buy_price"], values["sell_price"],
                     values["buy_date"], values["sell_date"], values["commission_override"], values["notes"], tid),
                )
                sync_trade_profits_to_funds(get_db(), cid)
                sync_client_holdings_from_trades(cid, get_db())
                invalidate_stock_caches()
                flash("Trade updated.", "success")
            else:
                insert_trade(cid, values)
                sync_trade_profits_to_funds(get_db(), cid)
                sync_client_holdings_from_trades(cid, get_db())
                invalidate_stock_caches()
                flash("Trade added.", "success")
            if request.form.get("add_another"):
                return redirect(url_for("trade_form", client=name))
            return redirect(safe_redirect_target(request.form.get("back"), url_for("trades_list")))
        clients = get_db().execute("SELECT name FROM clients ORDER BY lower(name)").fetchall()
        return render_template(
            "trade_form.html", trade=trade, clients=clients,
            preset_client=request.args.get("client", ""),
            back=safe_redirect_target(request.referrer, url_for("trades_list")),
        )

    @app.route("/trades/<int:tid>/delete", methods=["POST"])
    @login_required
    def trade_delete(tid):
        trade = get_db().execute("SELECT * FROM trades WHERE id=?", (tid,)).fetchone()
        if not trade:
            abort(404)
        get_db().execute("DELETE FROM trades WHERE id=?", (tid,))
        get_db().execute("DELETE FROM funds WHERE client_id=? AND notes LIKE ?", (trade["client_id"], f"%[Trade #{tid}]%"))
        flash("Trade deleted.", "success")
        return redirect(safe_redirect_target(request.referrer, url_for("trades_list")))

    @app.route("/trades/<int:tid>/delete-screenshot", methods=["POST"])
    @login_required
    def trade_delete_screenshot(tid):
        db = get_db()
        trade = db.execute("SELECT * FROM trades WHERE id=?", (tid,)).fetchone()
        if not trade:
            abort(404)
        upload_id = trade["upload_id"]
        if not upload_id:
            flash("No screenshot attached to this trade.", "info")
            return redirect(safe_redirect_target(request.referrer, url_for("trades_list")))

        stock_name = trade["stock"]
        db.execute("UPDATE trades SET upload_id=NULL WHERE id=?", (tid,))

        other_trades = db.execute("SELECT COUNT(*) AS c FROM trades WHERE upload_id=?", (upload_id,)).fetchone()
        count = other_trades["c"] if other_trades else 0
        if count == 0:
            try:
                assets = db.execute("SELECT path FROM upload_assets WHERE upload_id=?", (upload_id,)).fetchall()
                paths = [a["path"] for a in assets if a and a["path"]]
                if paths:
                    try:
                        storage = get_storage()
                        storage.remove(paths)
                    except Exception as storage_exc:
                        current_app.logger.warning("Failed to remove storage files for upload %s: %s", upload_id, storage_exc)
                db.execute("DELETE FROM upload_assets WHERE upload_id=?", (upload_id,))
                db.execute("DELETE FROM uploads WHERE id=?", (upload_id,))
            except Exception as exc:
                current_app.logger.warning("Error cleaning up orphaned upload record %s: %s", upload_id, exc)

        db.commit()
        flash(f"Proof screenshot removed from {stock_name}.", "success")
        return redirect(safe_redirect_target(request.referrer, url_for("trades_list")))

    @app.route("/trades/<int:tid>/close", methods=["POST"])
    @login_required
    def trade_close(tid):
        trade = get_db().execute("SELECT * FROM trades WHERE id=?", (tid,)).fetchone()
        if not trade:
            abort(404)
        if trade["buy_price"] is None:
            raise ValidationError("Add the opening buy price and date before closing this position.")
        if trade["sell_price"] is not None:
            abort(409, description="This position is already closed. Use Edit to correct it.")
        trade.update(
            sell_price=number(request.form.get("sell_price"), "Sell price", required=True),
            sell_date=valid_date(request.form.get("sell_date"), "Sell date", required=True),
        )
        values = trade_values(trade)
        get_db().execute(
            "UPDATE trades SET sell_price=?,sell_date=? WHERE id=?",
            (values["sell_price"], values["sell_date"], tid),
        )
        # Record realized profit to bank funds once
        pnl = ((dec(values["sell_price"]) - dec(trade["buy_price"])) * dec(trade["quantity"])).quantize(CENT, rounding=ROUND_HALF_UP)
        if pnl != ZERO:
            tag = f"[Trade #{tid}]"
            existing = get_db().execute(
                "SELECT id FROM funds WHERE client_id=? AND notes LIKE ?",
                (trade["client_id"], f"%{tag}%")
            ).fetchone()
            if not existing:
                pnl_note = f"[Realized {'Profit' if pnl > 0 else 'Loss'}] {trade['stock']} ({trade['quantity']} sh) {tag}"
                get_db().execute(
                    "INSERT INTO funds (client_id, amount, added_on, notes) VALUES (?, ?, ?, ?)",
                    (trade["client_id"], pnl, values["sell_date"], pnl_note),
                )
        flash("Position closed and profit added to bank funds.", "success")
        return redirect(safe_redirect_target(request.referrer, url_for("trades_list")))

    @app.route("/trades/bulk-delete", methods=["POST"])
    @login_required
    def trades_bulk_delete():
        ids = list(dict.fromkeys(identifier(value, "Trade") for value in request.form.getlist("ids")))
        if not ids:
            raise ValidationError("Choose at least one trade to delete.")
        if len(ids) > MAX_ROWS:
            raise ValidationError(f"Delete at most {MAX_ROWS} trades at once.")
        for tid in ids:
            get_db().execute("DELETE FROM funds WHERE notes LIKE ?", (f"%[Trade #{tid}]%",))
        count = get_db().execute(
            f"DELETE FROM trades WHERE id IN ({','.join('?' for _ in ids)})", ids
        ).rowcount
        flash(f"Deleted {count} trade(s).", "success")
        return redirect(safe_redirect_target(request.referrer, url_for("trades_list")))

    @app.route("/export/trades.xlsx")
    @login_required
    def export_trades():
        args = {key: request.args.get(key, "") for key in ("client_id", "stock", "from", "to", "status")}
        trades = query_trades(args)
        workbook = Workbook()
        workbook.active.title = "Trades"
        sheet(workbook.active, [title for title, _ in EXPORT_COLS],
              [[trade[key] for _, key in EXPORT_COLS] for trade in trades])
        selected = identifier(args["client_id"]) if args["client_id"] else None
        clients = [get_client(selected)] if selected else get_db().execute(
            "SELECT * FROM clients ORDER BY lower(name)"
        ).fetchall()
        all_trades = query_trades({"client_id": selected}) if selected else query_trades({})
        paid, rows = payments_by_client(), []
        for client in clients:
            cid = client["id"]
            summary = summarize([trade for trade in trades if trade["client_id"] == cid])
            all_earned = sum((trade["commission"] for trade in all_trades if trade["client_id"] == cid), ZERO)
            received = paid.get(cid, ZERO)
            rows.append([
                client["name"], summary["count"], summary["closed"], summary["open"], summary["pnl"],
                summary["commission"], all_earned, received, all_earned - received,
            ])
        sheet(
            workbook.create_sheet("Summary"),
            ["Username", "Trades (filtered)", "Closed (filtered)", "Open (filtered)", "P&L (filtered)",
             "Commission (filtered)", "Commission (all time)", "Paid (all time)", "Balance (all time)"],
            rows,
        )
        sql = "SELECT c.name,p.amount,p.paid_on,p.mode,p.notes FROM payments p JOIN clients c ON c.id=p.client_id"
        payments = get_db().execute(
            sql + (" WHERE p.client_id=?" if selected else "") + " ORDER BY p.paid_on DESC,p.id DESC",
            (selected,) if selected else (),
        ).fetchall()
        sheet(
            workbook.create_sheet("Payments"), ["Username", "Amount (all time)", "Paid On", "Mode", "Notes"],
            [[payment[key] for key in ("name", "amount", "paid_on", "mode", "notes")] for payment in payments],
        )
        return workbook_response(workbook, f"trades_{date.today().isoformat()}.xlsx")

    @app.route("/export/template.xlsx")
    @login_required
    def export_template():
        workbook = Workbook()
        sheet(
            workbook.active,
            ["Username", "Stock", "Quantity", "Buy Price", "Sell Price", "Buy Date", "Sell Date", "Commission", "Notes"],
            [["Ravi", "TATAMOTORS", 10, 950.25, 1010, "2024-09-12", "2024-09-20", "", "optional"]],
        )
        return workbook_response(workbook, "import_template.xlsx")

    @app.route("/import", methods=["POST"])
    @login_required
    def import_trades():
        uploaded = request.files.get("file")
        if not uploaded or not uploaded.filename.lower().endswith(".xlsx"):
            raise ValidationError("Choose an .xlsx file.")
        try:
            with ZipFile(uploaded.stream) as archive:
                if len(archive.infolist()) > 5000 or sum(item.file_size for item in archive.infolist()) > 32 * 1024 * 1024:
                    raise ValidationError("This workbook is too large after decompression.")
            uploaded.stream.seek(0)
            workbook = load_workbook(uploaded.stream, read_only=True, data_only=False)
            try:
                worksheet = workbook.active
                if worksheet is None:
                    raise ValidationError("The workbook has no active worksheet.")
                if worksheet.max_column and worksheet.max_column > 100:
                    raise ValidationError("The import worksheet must contain at most 100 columns.")
                rows = []
                for row in worksheet.iter_rows():
                    if len(rows) > MAX_ROWS:
                        raise ValidationError(f"Import at most {MAX_ROWS} rows per workbook.")
                    if any(cell.data_type == "f" for cell in row):
                        raise ValidationError("Formula cells are not accepted. Paste their values before importing.")
                    rows.append([cell.value for cell in row])
            finally:
                workbook.close()
        except ValidationError:
            raise
        except (BadZipFile, InvalidFileException, ParseError, KeyError, OSError, ValueError) as exc:
            raise ValidationError("Could not read that Excel workbook.") from exc
        if len(rows) < 2:
            raise ValidationError("The workbook contains no trade rows.")
        headers = [str(value or "").strip().lower() for value in rows[0]]
        aliases = {
            "client_name": ("username", "user", "client", "name"),
            "stock": ("stock", "stock name", "symbol", "scrip"), "quantity": ("quantity", "qty"),
            "buy_price": ("buy price", "buy", "buy avg"), "sell_price": ("sell price", "sell", "sell avg"),
            "buy_date": ("buy date",), "sell_date": ("sell date",),
            "commission_override": ("commission",), "notes": ("notes", "remarks"),
        }
        indexes = {key: next((headers.index(alias) for alias in names if alias in headers), None)
                   for key, names in aliases.items()}
        if any(indexes[key] is None for key in ("client_name", "stock", "quantity")):
            raise ValidationError("The sheet needs Username, Stock, and Quantity columns.")
        validated = []
        for index, row in enumerate(rows[1:], 2):
            if all(value is None or value == "" for value in row):
                continue
            values = {key: row[position] if position is not None and position < len(row) else None
                      for key, position in indexes.items()}
            try:
                validated.append((client_name(values["client_name"]), trade_values(values)))
            except ValidationError as exc:
                raise ValidationError(f"Row {index}: {exc} No rows were imported.") from exc
        if not validated:
            raise ValidationError("The workbook contains no populated trade rows.")
        grouped = {}
        for name, values in validated:
            grouped.setdefault(name, []).append(values)
        for name, values in grouped.items():
            cid = get_or_create_client(name)
            insert_trades(cid, values, source="import")
            sync_trade_profits_to_funds(get_db(), cid)
            sync_client_holdings_from_trades(cid, get_db())
        invalidate_stock_caches()
        flash(f"Imported {len(validated)} trade(s).", "success")
        return redirect(url_for("trades_list"))

    @app.route("/holdings")
    @login_required
    def holdings_list():
        stock_query = request.args.get("stock", "").strip().upper()
        db = get_db()

        base_query = """
            SELECT h.id, h.stock, c.name as client_name, c.id as client_id, h.quantity, h.buy_price, h.added_on
            FROM client_shares h 
            JOIN clients c ON c.id = h.client_id 
            WHERE h.quantity > 0
        """

        if stock_query:
            raw_holdings = db.execute(
                base_query + " AND UPPER(h.stock) LIKE ? ORDER BY h.stock, lower(c.name)", 
                (f"%{stock_query}%",)
            ).fetchall()
        else:
            raw_holdings = db.execute(
                base_query + " ORDER BY h.stock, lower(c.name)"
            ).fetchall()

        prices = get_cached_prices(db)
        holdings = []
        total_invested = ZERO
        total_current_value = ZERO

        for h in raw_holdings:
            item = dict(h)
            qty = dec(item.get("quantity"))
            buy_p = dec(item.get("buy_price"))
            cost = qty * buy_p
            total_invested += cost
            item["quantity"] = int(round(qty))

            stock_sym = (item.get("stock") or "").strip().upper()
            quote = prices.get(stock_sym)
            if quote:
                cmp = quote["current_price"]
                cur_val = qty * cmp
                unrealized = cur_val - cost
                unrealized_pct = (unrealized / cost * 100) if cost else ZERO
                item.update(
                    cmp=cmp,
                    current_value=cur_val,
                    unrealized_pnl=unrealized,
                    unrealized_pct=unrealized_pct,
                    price_change=quote.get("change", ZERO),
                    price_change_pct=quote.get("change_pct", ZERO),
                    price_updated_at=quote.get("updated_at"),
                )
                total_current_value += cur_val
            else:
                item.update(
                    cmp=None,
                    current_value=cost,
                    unrealized_pnl=ZERO,
                    unrealized_pct=ZERO,
                    price_change=ZERO,
                    price_change_pct=ZERO,
                    price_updated_at=None,
                )
                total_current_value += cost
            holdings.append(item)

        total_unrealized_pnl = total_current_value - total_invested
        total_unrealized_pct = (total_unrealized_pnl / total_invested * 100) if total_invested else ZERO

        return render_template(
            "holdings.html",
            holdings=holdings,
            query=stock_query,
            total_invested=total_invested,
            total_current_value=total_current_value,
            total_unrealized_pnl=total_unrealized_pnl,
            total_unrealized_pct=total_unrealized_pct,
            known_stocks=get_known_stocks(get_db()),
            client_entered_stocks=get_client_entered_stocks(get_db()),
        )

    @app.route("/admin/update-prices", methods=["GET", "POST"])
    @app.route("/update-prices", methods=["GET", "POST"])
    @login_required
    def admin_update_prices():
        try:
            res = update_all_stock_prices(get_db())
            count = len(res.get("updated", []))
            failed = len(res.get("failed", []))
            if count:
                flash(f"Updated market prices for {count} stock(s).", "success")
            elif failed:
                flash(f"Could not fetch prices for {failed} stock(s). Check symbols or internet connection.", "warning")
            else:
                flash("No holdings currently recorded to update.", "info")
        except Exception as e:
            app.logger.exception("Error updating stock prices: %s", e)
            flash(f"Could not update market prices: {e}", "warning")
        return redirect(url_for("holdings_list"))

    @app.route("/admin/stocks/alias", methods=["POST"])
    @login_required
    def admin_save_stock_alias():
        canonical_raw = request.form.get("canonical_name", "").strip().upper()
        aliases_raw = request.form.get("aliases", "").strip()
        redirect_to = request.form.get("redirect_to") or url_for("holdings_list")

        if not canonical_raw:
            flash("Single standard term (Canonical name) is required.", "error")
            return redirect(redirect_to)

        db = get_db()
        ensure_stock_aliases_table(db)

        # Collect aliases from both selected client stocks and manual input (excluding canonical name itself)
        selected_aliases = [s.strip().upper() for s in request.form.getlist("selected_aliases") if s.strip()]
        manual_tokens = [t.strip().upper() for t in aliases_raw.split(",") if t.strip()]
        alias_tokens = list(set([a for a in (selected_aliases + manual_tokens) if a != canonical_raw]))

        # If canonical_raw was previously mapped as an alias of another stock, reassign it to be its own canonical
        try:
            db.execute("UPDATE stock_aliases SET canonical_name=? WHERE alias=?", (canonical_raw, canonical_raw))
        except Exception:
            pass

        # If any selected alias was previously a canonical stock, merge all its child aliases under canonical_raw
        for a in alias_tokens:
            try:
                db.execute("UPDATE stock_aliases SET canonical_name=? WHERE canonical_name=?", (canonical_raw, a))
            except Exception:
                pass

        # Clean up any aliases previously mapped under canonical_raw that were unselected by admin
        try:
            existing_rows = db.execute("SELECT alias FROM stock_aliases WHERE canonical_name=?", (canonical_raw,)).fetchall()
            for row in existing_rows:
                ea = row["alias"].strip().upper()
                if ea != canonical_raw and ea not in alias_tokens:
                    db.execute("DELETE FROM stock_aliases WHERE canonical_name=? AND alias=?", (canonical_raw, ea))
        except Exception:
            pass

        # Update all client portals: trades, client_shares, client_alerts, stock_prices to the main canonical name
        updated_trades = 0
        updated_shares = 0
        for a in alias_tokens:
            r1 = db.execute("UPDATE trades SET stock=? WHERE UPPER(TRIM(stock))=?", (canonical_raw, a))
            r2 = db.execute("UPDATE client_shares SET stock=? WHERE UPPER(TRIM(stock))=?", (canonical_raw, a))
            db.execute("UPDATE client_alerts SET stock=? WHERE UPPER(TRIM(stock))=?", (canonical_raw, a))
            try:
                db.execute(
                    "UPDATE stock_prices SET stock=? WHERE UPPER(TRIM(stock))=? AND NOT EXISTS (SELECT 1 FROM stock_prices sp2 WHERE sp2.stock=?)",
                    (canonical_raw, a, canonical_raw)
                )
                db.execute("DELETE FROM stock_prices WHERE UPPER(TRIM(stock))=?", (a,))
            except Exception:
                pass
            if hasattr(r1, "rowcount") and r1.rowcount:
                updated_trades += r1.rowcount
            if hasattr(r2, "rowcount") and r2.rowcount:
                updated_shares += r2.rowcount

        # Remove all other aliases so only the main standard name stays in stock_aliases
        try:
            db.execute("DELETE FROM stock_aliases WHERE canonical_name=? AND UPPER(TRIM(alias)) != ?", (canonical_raw, canonical_raw))
            for a in alias_tokens:
                db.execute("DELETE FROM stock_aliases WHERE UPPER(TRIM(alias))=? AND UPPER(TRIM(canonical_name)) != ?", (a, canonical_raw))
            db.execute(
                """INSERT INTO stock_aliases (canonical_name, alias) VALUES (?, ?)
                   ON CONFLICT(alias) DO UPDATE SET canonical_name=excluded.canonical_name""",
                (canonical_raw, canonical_raw)
            )
        except Exception as exc:
            current_app.logger.warning("Error saving stock alias: %s", exc)

        invalidate_stock_caches()
        flash(
            f"Configured main stock name '{canonical_raw}'. Updated {updated_trades} trade(s) and {updated_shares} position(s). All other aliases removed so only the main name stays.",
            "success"
        )
        return redirect(redirect_to)

    @app.route("/admin/stocks/alias/delete", methods=["POST"])
    @login_required
    def admin_delete_stock_alias():
        alias_to_delete = request.form.get("alias", "").strip().upper()
        canonical_to_delete = request.form.get("canonical_name", "").strip().upper()
        redirect_to = request.form.get("redirect_to") or url_for("holdings_list")

        db = get_db()
        ensure_stock_aliases_table(db)

        if alias_to_delete:
            db.execute("DELETE FROM stock_aliases WHERE UPPER(TRIM(alias))=?", (alias_to_delete,))
            invalidate_stock_caches()
            flash(f"Alias '{alias_to_delete}' removed.", "success")
        elif canonical_to_delete:
            db.execute("DELETE FROM stock_aliases WHERE UPPER(TRIM(canonical_name))=?", (canonical_to_delete,))
            invalidate_stock_caches()
            flash(f"Stock mapping for '{canonical_to_delete}' removed.", "success")

        return redirect(redirect_to)

    @app.route("/clients", methods=["GET", "POST"])
    @login_required
    def clients_list():
        if request.method == "POST":
            name = client_name(request.form.get("name"))
            values = client_values(request.form, name)
            if get_db().execute("SELECT id FROM clients WHERE lower(name)=lower(?)", (name,)).fetchone():
                raise ValidationError("A user with that name already exists.")
            raw_username = request.form.get("username", "").strip()
            username = raw_username.lstrip("@").strip() if raw_username else None
            if username:
                if len(username) < 3 or len(username) > 60:
                    raise ValidationError("Username must be between 3 and 60 characters.")
                if get_db().execute("SELECT id FROM clients WHERE lower(username)=lower(?)", (username,)).fetchone():
                    raise ValidationError("A client with that username already exists.")
            portal_pwd = (request.form.get("portal_password") or request.form.get("client_password") or "").strip()
            if not portal_pwd:
                portal_pwd = "Client@1234"
            p_hash = generate_password_hash(portal_pwd)
            p_plain = portal_pwd
            get_db().execute(
                "INSERT INTO clients(name,phone,email,commission_type,commission_rate,notes,portfolio_amount,password_hash,password_plain,username,password_changed) VALUES (?,?,?,?,?,?,?,?,?,?,0)",
                (*values, p_hash, p_plain, username),
            )
            flash(f"Added {name}." + (f" Username: @{username}." if username else "") + (f" Portal password set: {portal_pwd}" if portal_pwd else ""), "success")
            return redirect(url_for("clients_list"))
        trades, paid, rows = query_trades({}), payments_by_client(), []
        for client in get_db().execute("SELECT * FROM clients ORDER BY lower(name)").fetchall():
            summary = summarize([trade for trade in trades if trade["client_id"] == client["id"]])
            received = paid.get(client["id"], ZERO)
            rows.append({
                **client, **summary, "paid": received, "outstanding": summary["commission"] - received,
                "rule": rule_label(client["commission_type"], client["commission_rate"]),
            })
        return render_template("clients.html", clients=rows)

    @app.route("/clients/<int:cid>", methods=["GET", "POST"])
    @login_required
    def client_detail(cid):
        client = get_client(cid)
        if request.method == "POST":
            name = client_name(request.form.get("name"))
            values = client_values(request.form, name)
            if get_db().execute("SELECT id FROM clients WHERE lower(name)=lower(?) AND id<>?", (name, cid)).fetchone():
                raise ValidationError("Another user already has that name.")
            raw_username = request.form.get("username", "").strip()
            username = raw_username.lstrip("@").strip() if raw_username else None
            if username:
                if len(username) < 3 or len(username) > 60:
                    raise ValidationError("Username must be between 3 and 60 characters.")
                if get_db().execute("SELECT id FROM clients WHERE lower(username)=lower(?) AND id<>?", (username, cid)).fetchone():
                    raise ValidationError("Another client already has that username.")
            get_db().execute(
                "UPDATE clients SET name=?,phone=?,email=?,commission_type=?,commission_rate=?,notes=?,portfolio_amount=?,username=? WHERE id=?",
                (*values, username, cid),
            )
            global _active_clients_cache
            _active_clients_cache = None
            flash("User updated.", "success")
            return redirect(url_for("client_detail", cid=cid))
        trades = query_trades({"client_id": cid})
        summary = summarize(trades)
        payments = get_db().execute(
            "SELECT * FROM payments WHERE client_id=? ORDER BY paid_on DESC,id DESC", (cid,)
        ).fetchall()
        summary["paid"] = sum((dec(payment["amount"]) for payment in payments), ZERO)
        summary["outstanding"] = summary["commission"] - summary["paid"]

        try:
            funds = get_db().execute(
                "SELECT * FROM funds WHERE client_id=? ORDER BY added_on DESC,id DESC", (cid,)
            ).fetchall()
        except Exception:
            funds = []
        total_funds = sum((dec(fund["amount"]) for fund in funds), ZERO)

        raw_shares = get_client_calculated_holdings(cid, get_db(), sync_trades=False)
        prices = get_cached_prices(get_db())
        shares = []
        total_holdings_invested = ZERO
        for s in raw_shares:
            item = dict(s)
            qty = dec(item.get("quantity"))
            buy_p = dec(item.get("buy_price"))
            cost = qty * buy_p
            total_holdings_invested += cost
            quote = prices.get((item.get("stock") or "").strip().upper())
            if quote:
                cmp = quote["current_price"]
                cur_val = qty * cmp
                unrealized = cur_val - cost
                unrealized_pct = (unrealized / cost * 100) if cost else ZERO
                item.update(
                    cmp=cmp,
                    current_value=cur_val,
                    unrealized_pnl=unrealized,
                    unrealized_pct=unrealized_pct,
                    price_change=quote.get("change", ZERO),
                    price_change_pct=quote.get("change_pct", ZERO),
                )
            else:
                item.update(
                    cmp=None,
                    current_value=cost,
                    unrealized_pnl=ZERO,
                    unrealized_pct=ZERO,
                    price_change=ZERO,
                    price_change_pct=ZERO,
                )
            shares.append(item)

        summary["money_in_bank"] = total_funds - total_holdings_invested

        try:
            raw_advisories = get_db().execute(
                """
                SELECT a.*, v.acknowledged, v.acknowledged_at
                FROM client_alerts a
                LEFT JOIN client_alert_views v ON v.alert_id = a.id AND v.client_id = ?
                WHERE a.client_id = ? OR a.client_id IS NULL
                ORDER BY a.id DESC
                """,
                (cid, cid),
            ).fetchall()
            alias_map = get_alias_mapping(get_db())
            held_canonical_stocks = {
                alias_map.get((sh.get("stock") or "").strip().upper(), (sh.get("stock") or "").strip().upper())
                for sh in shares
                if dec(sh.get("quantity") or 0) > ZERO
            }
            advisories = []
            for adv in raw_advisories:
                if adv["action"] == "SELL" and adv["client_id"] is None:
                    adv_stock = alias_map.get((adv.get("stock") or "").strip().upper(), (adv.get("stock") or "").strip().upper())
                    if adv_stock not in held_canonical_stocks and not adv.get("acknowledged"):
                        continue
                advisories.append(adv)
        except Exception:
            advisories = []

        return render_template(
            "client_detail.html", client=client, trades=trades, s=summary, payments=payments,
            funds=funds, total_funds=total_funds, shares=shares, advisories=advisories,
            uploads=upload_records(cid), rule=rule_label(client["commission_type"], client["commission_rate"]),
            merge_clients=get_db().execute("SELECT id,name FROM clients WHERE id<>? ORDER BY lower(name)", (cid,)).fetchall(),
            payment_token=payment_signer().dumps(cid),
        )

    @app.route("/clients/<int:cid>/set-username", methods=["POST"])
    @login_required
    def client_set_username(cid):
        client = get_client(cid)
        raw_username = request.form.get("username", "").strip()
        username = raw_username.lstrip("@").strip() if raw_username else None
        back = request.form.get("back")
        if username:
            if len(username) < 3 or len(username) > 60:
                flash("Username must be between 3 and 60 characters.", "error")
                return redirect(safe_redirect_target(back, url_for("client_detail", cid=cid)))
            if not re.match(r"^[a-zA-Z0-9_.]+$", username):
                flash("Username can only contain letters, numbers, underscores, and dots.", "error")
                return redirect(safe_redirect_target(back, url_for("client_detail", cid=cid)))
            existing = get_db().execute("SELECT id FROM clients WHERE lower(username)=lower(?) AND id<>?", (username, cid)).fetchone()
            if existing:
                flash("A client with that username already exists.", "error")
                return redirect(safe_redirect_target(back, url_for("client_detail", cid=cid)))
        get_db().execute("UPDATE clients SET username=? WHERE id=?", (username, cid))
        msg = f"Username for {client['name']} updated to: @{username}" if username else f"Username removed for {client['name']}"
        flash(msg, "success")
        return redirect(safe_redirect_target(back, url_for("client_detail", cid=cid)))

    @app.route("/clients/<int:cid>/set-password", methods=["POST"])
    @login_required
    def client_set_password(cid):
        client = get_client(cid)
        pwd = request.form.get("client_password", "").strip()
        confirm = request.form.get("confirm_password", "").strip()
        back = request.form.get("back")
        if not pwd or len(pwd) < 4:
            flash("Password must be at least 4 characters.", "error")
            return redirect(safe_redirect_target(back, url_for("client_detail", cid=cid)))
        if confirm and pwd != confirm:
            flash("Passwords do not match.", "error")
            return redirect(safe_redirect_target(back, url_for("client_detail", cid=cid)))
        p_hash = generate_password_hash(pwd)
        get_db().execute("UPDATE clients SET password_hash=?, password_plain=? WHERE id=?", (p_hash, pwd, cid))
        flash(f"Portal password for {client['name']} updated to: {pwd}", "success")
        return redirect(safe_redirect_target(back, url_for("client_detail", cid=cid)))

    @app.route("/clients/<int:cid>/shares/new", methods=["POST"])
    @login_required
    def client_share_add(cid):
        get_client(cid)
        db = get_db()
        stock = resolve_canonical_stock(db, request.form.get("stock", ""))
        quantity = number(request.form.get("quantity"), "Quantity") if request.form.get("quantity") else None
        buy_price = number(request.form.get("buy_price"), "Buy Price") if request.form.get("buy_price") else None
        if not stock:
            abort(400)
        qty_int = int(round(quantity)) if quantity is not None else None

        existing = db.execute(
            "SELECT id, quantity, buy_price FROM client_shares WHERE client_id=? AND UPPER(TRIM(stock))=?",
            (cid, stock.upper())
        ).fetchone()
        if existing:
            old_qty = dec(existing["quantity"] or 0)
            old_bp = dec(existing["buy_price"] or 0)
            add_qty = dec(qty_int or 0)
            add_bp = dec(buy_price or old_bp)
            new_qty = int(round(old_qty + add_qty))
            if new_qty > 0 and (old_qty * old_bp + add_qty * add_bp) > 0:
                new_avg = ((old_qty * old_bp + add_qty * add_bp) / dec(new_qty)).quantize(CENT, rounding=ROUND_HALF_UP)
            else:
                new_avg = add_bp
            db.execute(
                "UPDATE client_shares SET quantity=?, buy_price=? WHERE id=?",
                (new_qty, new_avg, existing["id"])
            )
        else:
            db.execute(
                "INSERT INTO client_shares (client_id, stock, quantity, buy_price) VALUES (?, ?, ?, ?)",
                (cid, stock.upper(), qty_int, buy_price)
            )
        db.commit()
        flash("Holding added.", "success")
        return redirect(safe_redirect_target(request.form.get("back"), url_for("client_detail", cid=cid)))

    @app.route("/clients/<int:cid>/shares/<int:sid>/delete", methods=["POST"])
    @login_required
    def client_share_delete(cid, sid):
        db = get_db()
        holding = db.execute("SELECT * FROM client_shares WHERE id=? AND client_id=?", (sid, cid)).fetchone()
        deleted = db.execute("DELETE FROM client_shares WHERE id=? AND client_id=?", (sid, cid)).rowcount
        if not deleted:
            trade = db.execute("SELECT * FROM trades WHERE id=? AND client_id=?", (sid, cid)).fetchone()
            if trade:
                db.execute("DELETE FROM trades WHERE id=? AND client_id=?", (sid, cid))
                db.execute("DELETE FROM client_shares WHERE client_id=? AND UPPER(TRIM(stock))=?", (cid, trade["stock"].upper()))
        else:
            if holding:
                try:
                    db.execute(
                        "DELETE FROM trades WHERE client_id=? AND UPPER(TRIM(stock))=? AND (sell_price IS NULL OR sell_price=0)",
                        (cid, holding["stock"].upper())
                    )
                except Exception:
                    pass
        db.commit()
        flash("Holding deleted.", "success")
        return redirect(safe_redirect_target(request.referrer, url_for("client_detail", cid=cid)))

    @app.route("/clients/<int:cid>/shares/<int:sid>/edit", methods=["POST"])
    @login_required
    def client_share_edit(cid, sid):
        db = get_db()
        stock = resolve_canonical_stock(db, request.form.get("stock", ""))
        quantity = number(request.form.get("quantity"), "Quantity") if request.form.get("quantity") else None
        buy_price = number(request.form.get("buy_price"), "Buy Price") if request.form.get("buy_price") else None
        if not stock:
            abort(400)
        qty_int = int(round(quantity)) if quantity is not None else None

        # 1. Update client_shares with id=sid
        res = db.execute(
            "UPDATE client_shares SET stock=?, quantity=?, buy_price=? WHERE id=? AND client_id=?",
            (stock.upper(), qty_int, buy_price, sid, cid)
        )

        # 2. If no rows updated (e.g. sid was a trade ID or missing from client_shares)
        if not getattr(res, "rowcount", 0):
            trade_row = db.execute("SELECT * FROM trades WHERE id=? AND client_id=?", (sid, cid)).fetchone()
            if trade_row:
                db.execute(
                    "UPDATE trades SET stock=?, quantity=?, buy_price=? WHERE id=? AND client_id=?",
                    (stock.upper(), qty_int, buy_price, sid, cid)
                )

            existing_cs = db.execute(
                "SELECT id FROM client_shares WHERE client_id=? AND UPPER(TRIM(stock))=?",
                (cid, stock.upper())
            ).fetchone()
            if existing_cs:
                db.execute(
                    "UPDATE client_shares SET quantity=?, buy_price=? WHERE id=?",
                    (qty_int, buy_price, existing_cs["id"])
                )
            else:
                db.execute(
                    "INSERT INTO client_shares (client_id, stock, quantity, buy_price) VALUES (?, ?, ?, ?)",
                    (cid, stock.upper(), qty_int, buy_price)
                )
        else:
            # Also keep any open trade for this stock in sync
            try:
                db.execute(
                    """UPDATE trades SET quantity=?, buy_price=?
                       WHERE client_id=? AND UPPER(TRIM(stock))=? AND buy_price IS NOT NULL AND (sell_price IS NULL OR sell_price=0)""",
                    (qty_int, buy_price, cid, stock.upper())
                )
            except Exception:
                pass

        db.commit()
        flash("Holding updated.", "success")
        return redirect(safe_redirect_target(request.referrer, url_for("client_detail", cid=cid)))

    @app.route("/clients/<int:cid>/shares/<int:sid>/sell", methods=["POST"])
    @login_required
    def client_share_sell(cid, sid):
        db = get_db()
        holding = db.execute("SELECT * FROM client_shares WHERE id=? AND client_id=?", (sid, cid)).fetchone()
        if not holding:
            trade_row = db.execute("SELECT * FROM trades WHERE id=? AND client_id=?", (sid, cid)).fetchone()
            if trade_row:
                holding = {
                    "id": sid,
                    "client_id": cid,
                    "stock": trade_row["stock"],
                    "quantity": trade_row["quantity"],
                    "buy_price": trade_row["buy_price"],
                    "added_on": trade_row["buy_date"],
                }
            else:
                abort(404)

        sell_qty = int(round(number(request.form.get("quantity"), "Quantity", required=True, positive=True)))
        sell_price = number(request.form.get("sell_price"), "Sell Price", required=True, positive=True)
        sell_date = valid_date(request.form.get("sell_date"), "Sell Date") or date.today().isoformat()

        new_qty = int(round(dec(holding["quantity"] or ZERO) - dec(sell_qty)))
        if new_qty <= 0:
            db.execute("DELETE FROM client_shares WHERE id=?", (holding["id"],))
            db.execute("DELETE FROM client_shares WHERE client_id=? AND UPPER(TRIM(stock))=?", (cid, holding["stock"].upper()))
        else:
            db.execute("UPDATE client_shares SET quantity=? WHERE id=?", (new_qty, holding["id"]))

        buy_date = holding["added_on"][:10] if holding.get("added_on") else date.today().isoformat()

        trade_ids = insert_trades(cid, [{
            "stock": holding["stock"],
            "quantity": sell_qty,
            "buy_price": holding["buy_price"],
            "sell_price": sell_price,
            "buy_date": buy_date,
            "sell_date": sell_date,
            "notes": "Sold from Holdings"
        }])
        trade_id = trade_ids[0] if trade_ids else None
        pnl = ((dec(sell_price) - dec(holding["buy_price"])) * dec(sell_qty)).quantize(CENT, rounding=ROUND_HALF_UP)
        if pnl != ZERO:
            tag = f"[Trade #{trade_id}]" if trade_id else f"[Holding Sale #{sid}]"
            existing = db.execute(
                "SELECT id FROM funds WHERE client_id=? AND notes LIKE ?",
                (cid, f"%{tag}%")
            ).fetchone()
            if not existing:
                pnl_note = f"[Realized {'Profit' if pnl > 0 else 'Loss'}] {holding['stock']} ({sell_qty} sh) {tag}"
                db.execute(
                    "INSERT INTO funds (client_id, amount, added_on, notes) VALUES (?, ?, ?, ?)",
                    (cid, pnl, sell_date, pnl_note),
                )

        db.commit()
        flash(f"Sold {sell_qty} shares of {holding['stock']} and recorded in Trade Ledger.", "success")
        return redirect(safe_redirect_target(request.referrer, url_for("client_detail", cid=cid)))

    @app.route("/clients/<int:cid>/delete", methods=["POST"])
    @login_required
    def client_delete(cid):
        get_client(cid)
        get_db().execute("DELETE FROM clients WHERE id=?", (cid,))
        flash("User and their trades/payments deleted. Submission evidence is retained.", "success")
        return redirect(url_for("clients_list"))

    @app.route("/clients/<int:cid>/merge", methods=["POST"])
    @login_required
    def client_merge(cid):
        get_client(cid)
        target = identifier(request.form.get("target_id"))
        get_client(target)
        if target == cid:
            raise ValidationError("Choose a different user to merge into.")
        for table in ("trades", "payments", "uploads"):
            get_db().execute(f"UPDATE {table} SET client_id=? WHERE client_id=?", (target, cid))
        get_db().execute("DELETE FROM clients WHERE id=?", (cid,))
        flash("Users merged, including their submission history.", "success")
        return redirect(url_for("client_detail", cid=target))

    @app.route("/payments", methods=["GET", "POST"])
    @login_required
    def payments_list():
        if request.method == "POST":
            cid = identifier(request.form.get("client_id"))
            amount = number(request.form.get("amount"), "Payment amount", required=True, positive=True, places=2)
            paid_on = valid_date(request.form.get("paid_on"), "Payment date", required=True)
            mode = text(request.form.get("mode"), "Payment mode", limit=80) or None
            notes = text(request.form.get("notes"), "Notes", multiline=True) or None
            get_client(cid)
            get_db().execute(
                "INSERT INTO payments(client_id,amount,paid_on,mode,notes) VALUES (?,?,?,?,?)",
                (cid, amount, paid_on, mode, notes),
            )
            flash("Payment recorded.", "success")
            return redirect(safe_redirect_target(request.form.get("back"), url_for("payments_list")))
        payments = get_db().execute(
            "SELECT p.*,c.name AS client_name FROM payments p JOIN clients c ON c.id=p.client_id "
            "ORDER BY p.paid_on DESC,p.id DESC"
        ).fetchall()
        clients = get_db().execute("SELECT id,name FROM clients ORDER BY lower(name)").fetchall()
        preset = identifier(request.args["client_id"]) if request.args.get("client_id") else None
        return render_template(
            "payments.html", payments=payments, clients=clients,
            total=sum((dec(payment["amount"]) for payment in payments), ZERO), preset=preset,
        )

    @app.route("/payments/<int:pid>/delete", methods=["POST"])
    @login_required
    def payment_delete(pid):
        if get_db().execute("DELETE FROM payments WHERE id=?", (pid,)).rowcount != 1:
            abort(404)
        flash("Payment deleted.", "success")
        return redirect(safe_redirect_target(request.referrer, url_for("payments_list")))

    @app.route("/funds", methods=["POST"])
    @login_required
    def add_fund():
        cid = identifier(request.form.get("client_id"))
        amount = number(request.form.get("amount"), "Fund amount", required=True, positive=True, places=2)
        added_on = valid_date(request.form.get("added_on"), "Date", required=True)
        notes = text(request.form.get("notes"), "Notes", multiline=True) or None
        get_client(cid)
        get_db().execute(
            "INSERT INTO funds(client_id,amount,added_on,notes) VALUES (?,?,?,?)",
            (cid, amount, added_on, notes),
        )
        flash("Funds added.", "success")
        return redirect(safe_redirect_target(request.form.get("back"), url_for("client_detail", cid=cid)))

    @app.route("/funds/<int:fid>/delete", methods=["POST"])
    @login_required
    def fund_delete(fid):
        if get_db().execute("DELETE FROM funds WHERE id=?", (fid,)).rowcount != 1:
            abort(404)
        flash("Funds deleted.", "success")
        return redirect(safe_redirect_target(request.referrer, url_for("dashboard")))

    @app.route("/uploads")
    @login_required
    def uploads_list():
        return render_template("uploads.html", uploads=upload_records())

    @app.route("/uploads/file/<path:name>")
    def upload_file(name):
        if not session.get("admin") and not session.get("client_id"):
            abort(401)
        safe_name = safe_path(name)
        if session.get("client_id") and not session.get("admin"):
            allowed = get_db().execute(
                "SELECT 1 FROM uploads WHERE filename LIKE ? AND client_id=?",
                (f"%{safe_name}%", session["client_id"]),
            ).fetchone()
            if not allowed:
                abort(403)
        if app.config["BACKEND_MODE"] == "supabase":
            asset = get_db().execute(
                "SELECT path FROM upload_assets WHERE path=? AND upload_id IS NOT NULL", (safe_name,)
            ).fetchone()
            if not asset:
                abort(404)
            return redirect(get_storage().sign_download(safe_name), code=302)
        return send_from_directory(app.config["UPLOAD_DIR"], safe_name, as_attachment=False)

    @app.route("/admin/scan")
    @login_required
    def admin_scan():
        owner_hash()
        return render_template("upload.html", admin=True, need_code=False, portal_token="")

    @app.route("/settings", methods=["GET", "POST"])
    @login_required
    def settings_page():
        if request.method == "POST":
            ctype, rate = rule_values(
                request.form.get("default_commission_type"), request.form.get("default_commission_rate")
            )
            values = {
                "default_commission_type": ctype, "default_commission_rate": str(rate),
                "upload_code": text(request.form.get("upload_code"), "Upload code", limit=128),
                "currency": text(request.form.get("currency"), "Currency", limit=12, required=True),
                "business_name": text(request.form.get("business_name"), "Business name", limit=120),
            }
            new_password = request.form.get("new_password", "")
            if new_password:
                if app.config["ADMIN_PASSWORD_HASH"]:
                    raise ValidationError("This password is managed on the server. Update ADMIN_PASSWORD or ADMIN_PASSWORD_HASH.")
                if not 10 <= len(new_password) <= 1024:
                    raise ValidationError("Password must contain between 10 and 1024 characters.")
                values["admin_password"] = generate_password_hash(new_password)
            for key, value in values.items():
                set_setting(key, value)
            if new_password:
                session["auth_version"] = admin_credentials()[1]
                session["_csrf"] = secrets.token_urlsafe(32)
            flash("Settings saved.", "success")
            return redirect(url_for("settings_page"))
        keys = ("default_commission_type", "default_commission_rate", "upload_code", "currency", "business_name")
        legacy_count = get_db().execute(
            "SELECT COUNT(*) AS total FROM trades WHERE commission_type_snapshot IS NULL"
        ).fetchone()["total"]
        base = app.config["PUBLIC_BASE_URL"] or request.host_url.rstrip("/")
        go_url = base + url_for("short_upload_link")
        short_url = setting("upload_short_url") or go_url
        return render_template(
            "settings.html", s={key: setting(key, "") for key in keys}, upload_url=portal_url(),
            short_url=short_url, legacy_commission_count=legacy_count,
            known_stocks=get_known_stocks(get_db()),
        )

    @app.route("/settings/rotate-link", methods=["POST"])
    @login_required
    def rotate_upload_link():
        new_token = secrets.token_urlsafe(32)
        set_setting("upload_token", new_token)
        g.pop("settings_values", None)
        # Regenerate the short URL for the new token
        try:
            short = shorten_url(portal_url())
            if short:
                set_setting("upload_short_url", short)
            else:
                set_setting("upload_short_url", "")
        except Exception:
            set_setting("upload_short_url", "")
        flash("A new submission link is ready. The old link and its pending image authorizations no longer work.", "success")
        return redirect(url_for("settings_page"))

    @app.route("/go")
    def short_upload_link():
        token = setting("upload_token", "")
        if not token:
            abort(404, description="No submission link has been set up yet.")
        return redirect(url_for("public_upload", token=token), 302)

    @app.route("/upload", defaults={"token": None})
    @app.route("/submit/<token>")
    def public_upload(token=None):
        if token is None:
            if session.get("admin"):
                return redirect(url_for("public_upload", token=portal_token()))
            return redirect(url_for("login"))
        expected = setting("upload_token", "")
        if not expected or not hmac.compare_digest(token.encode(), expected.encode()):
            abort(404, description="This submission link is inactive. Ask the administrator for the current link.")
        owner_hash()
        return render_template(
            "upload.html", admin=False, need_code=bool(setting("upload_code")), portal_token=token,
            og_title="Commission Book – Send Your Trades",
            og_description="Submit your trades in seconds. Just fill in the share name and amount — no account needed.",
        )

    @app.route("/pay/<token>", methods=["GET", "POST"])
    def client_pay(token):
        try:
            cid = payment_signer().loads(token, max_age=30*24*3600)
        except Exception:
            abort(404, description="Invalid or expired payment link.")
            
        client = get_client(cid)
        trades = query_trades({"client_id": cid})
        summary = summarize(trades)
        paid = sum((dec(payment["amount"]) for payment in get_db().execute("SELECT amount FROM payments WHERE client_id=?", (cid,)).fetchall()), ZERO)
        outstanding = summary["commission"] - paid
        summary["paid"] = paid
        
        if request.method == "POST":
            if outstanding > 0:
                get_db().execute(
                    "INSERT INTO payments(client_id,amount,paid_on,mode,notes) VALUES (?,?,?,?,?)",
                    (cid, outstanding, date.today().isoformat(), "Client Link", "Marked as paid via client link"),
                )
                
            import smtplib
            from email.message import EmailMessage
            admin_email = os.environ.get("ADMIN_EMAIL")
            if admin_email:
                try:
                    msg = EmailMessage()
                    msg.set_content(f"Client {client['name']} marked their outstanding balance of {outstanding} as paid.\nCheck progress at: {url_for('client_detail', cid=client['id'], _external=True)}")
                    msg['Subject'] = f"Payment marked as done by {client['name']}"
                    msg['From'] = admin_email
                    msg['To'] = admin_email
                    
                    s = smtplib.SMTP(os.environ.get("SMTP_SERVER", "smtp.gmail.com"), int(os.environ.get("SMTP_PORT", 587)))
                    s.starttls()
                    s.login(os.environ.get("SMTP_USER", admin_email), os.environ.get("SMTP_PASSWORD", ""))
                    s.send_message(msg)
                    s.quit()
                except Exception as e:
                    app.logger.error(f"Failed to send email: {e}")
            
            flash("Thank you, payment marked as done.", "success")
            return redirect(url_for("client_pay", token=token))
            
        return render_template("pay.html", client=client, outstanding=outstanding, s=summary, trades=trades)

    @app.route("/api/parse", methods=["POST"])
    def api_parse():
        data = json_object()
        authorize_upload(data)
        if not isinstance(data.get("text"), str):
            raise ValidationError("OCR text must be a string.")
        content = text(data["text"], "OCR text", limit=MAX_TEXT, multiline=True)
        return jsonify(parse_text(content))

    @app.route("/api/uploads/sign", methods=["POST"])
    def sign_upload():
        data = json_object()
        authorization = authorize_upload(data)
        text(data.get("filename"), "Filename", limit=255, required=True)
        content_type = data.get("content_type")
        if not isinstance(content_type, str) or content_type not in IMAGE_TYPES:
            raise ValidationError("Only PNG, JPEG, and WebP screenshots are supported.")
        maximum = MAX_CLOUD_IMAGE if app.config["BACKEND_MODE"] == "supabase" else MAX_LOCAL_IMAGE
        size = data.get("size")
        if isinstance(size, bool) or not isinstance(size, int) or not 0 < size <= maximum:
            raise ValidationError(f"Each image must contain between 1 and {maximum:,} bytes.")
        path = "evidence/" + uuid.uuid4().hex + IMAGE_TYPES[content_type][1]
        receipt = receipt_signer().dumps({**authorization, "path": path})
        if app.config["BACKEND_MODE"] == "supabase":
            storage = get_storage()
            bucket = storage.check_bucket()
            maximum = min(maximum, bucket["file_size_limit"])
            if size > maximum:
                raise ValidationError(f"This private bucket accepts images up to {maximum:,} bytes.")
            result = storage.sign_upload(path, content_type)
        else:
            result = {
                "upload_url": url_for("local_upload", receipt=receipt, _external=True), "method": "PUT",
                "headers": {"Content-Type": content_type, "X-CSRF-Token": csrf_token()},
            }
        get_db().execute(
            """INSERT INTO upload_assets(path,owner_hash,scope_hash,content_type,byte_size,expires_at)
            VALUES (?,?,?,?,?,?)""",
            (path, authorization["owner"], authorization["scope"], content_type, size, int(time.time()) + UPLOAD_TTL),
        )
        return jsonify(
            **result, path=path, receipt=receipt, expires_in=UPLOAD_TTL, max_size=maximum,
        )

    @app.route("/api/uploads/local/<receipt>", methods=["PUT"])
    def local_upload(receipt):
        if app.config["BACKEND_MODE"] != "local":
            abort(404)
        decoded = decode_receipt(receipt)
        asset = receipt_asset(decoded.get("path"), receipt)
        if request.mimetype != asset["content_type"]:
            raise ValidationError("The image Content-Type does not match its authorization.")
        contents = request.get_data()
        if len(contents) != asset["byte_size"]:
            raise ValidationError("The image size does not match its authorization.")
        get_storage().store(asset["path"], contents, asset["content_type"])
        return "", 204

    @app.route("/api/submit", methods=["POST"])
    def api_submit():
        data = json_object()
        authorization = authorize_upload(data)
        name = client_name(data.get("name"))
        raw_phone = data.get("phone")
        phone = clean_phone_number(raw_phone)
        if not phone or len(phone) != 10:
            raise ValidationError("Please provide a valid 10-digit mobile number.")
        submission_id = submission_uuid(data.get("submission_id"))
        rows = data.get("rows")
        if not isinstance(rows, list) or not 1 <= len(rows) <= MAX_ROWS:
            raise ValidationError(f"Submit between 1 and {MAX_ROWS} trade rows.")
        validated = []
        for index, row in enumerate(rows, 1):
            try:
                validated.append(trade_values(row, public=True))
            except ValidationError as exc:
                raise ValidationError(f"Row {index}: {exc} No trades were saved.") from exc
        images = data.get("images", [])
        if not isinstance(images, list) or len(images) > 10:
            raise ValidationError("Attach at most 10 screenshots.")
        paths = []
        for image in images:
            if not isinstance(image, dict) or set(image) != {"path", "receipt"}:
                raise ValidationError("Every image must contain its path and signed receipt.")
            paths.append(safe_path(image["path"]))
            if not isinstance(image["receipt"], str) or len(image["receipt"]) > 2048:
                raise ValidationError("Invalid image receipt.")
        if len(set(paths)) != len(paths):
            raise ValidationError("An image cannot be attached more than once.")
        ocr_text = text(data.get("ocr_text"), "OCR text", limit=MAX_TEXT, multiline=True)
        platform = text(data.get("platform"), "Platform", limit=80) or None
        normalized = {
            "name": name, "phone": phone, "rows": validated, "images": images, "ocr_text": ocr_text,
            "platform": platform, "scope": authorization["scope"],
        }
        digest = hashlib.sha256(json.dumps(
            normalized, sort_keys=True, separators=(",", ":"), default=str,
        ).encode()).hexdigest()
        previous = prior_submission(submission_id, digest, authorization["owner"])
        if previous == "CONFLICT":
            submission_id = str(uuid.uuid4())
            previous = None
        if previous is not None:
            return previous
        assets = [receipt_asset(image["path"], image["receipt"], authorization) for image in images]
        storage = get_storage()
        for asset in assets:
            storage.verify(asset)
        cid = get_or_create_client(name, phone)
        row = get_db().execute(
            """INSERT INTO uploads(
                client_name,client_id,filename,platform,ocr_text,trade_count,
                submission_id,submission_hash,submission_owner
            ) VALUES (?,?,?,?,?,?,?,?,?) ON CONFLICT(submission_id) DO NOTHING RETURNING id""",
            (name, cid, ",".join(paths), platform, ocr_text, len(validated),
             submission_id, digest, authorization["owner"]),
        ).fetchone()
        if row is None:
            previous = prior_submission(submission_id, digest, authorization["owner"])
            if previous == "CONFLICT":
                submission_id = str(uuid.uuid4())
                previous = None
            if previous is not None:
                return previous
            # If we still hit a conflict (unlikely), generate a new ID and retry the insert
            submission_id = str(uuid.uuid4())
            row = get_db().execute(
                """INSERT INTO uploads(
                    client_name,client_id,filename,platform,ocr_text,trade_count,
                    submission_id,submission_hash,submission_owner
                ) VALUES (?,?,?,?,?,?,?,?,?) ON CONFLICT(submission_id) DO NOTHING RETURNING id""",
                (name, cid, ",".join(paths), platform, ocr_text, len(validated),
                 submission_id, digest, authorization["owner"]),
            ).fetchone()
        upload_id = row["id"]
        for asset in assets:
            get_db().execute(
                "UPDATE upload_assets SET upload_id=? WHERE path=? AND upload_id IS NULL",
                (upload_id, asset["path"]),
            )
        insert_trades(cid, validated, source="upload", upload_id=upload_id)
        sync_trade_profits_to_funds(get_db(), cid)
        sync_client_holdings_from_trades(cid, get_db())
        invalidate_stock_caches()
        return jsonify(ok=True, saved=len(validated), submission_id=submission_id)

    @app.cli.command("init-db")
    def init_database():
        if app.config["BACKEND_MODE"] != "local":
            raise click.ClickException("Run supabase/migrations/001_commission_book.sql in Supabase instead.")
        get_db()
        click.echo("Local schema is ready. Existing ledger entries have been preserved.")

    @app.cli.command("check-cloud")
    def check_cloud():
        if app.config["BACKEND_MODE"] != "supabase":
            raise click.ClickException("DATABASE_URL is not configured; this application is using local SQLite.")
        try:
            database = get_db()
            database.execute("SELECT submission_id,client_id FROM uploads LIMIT 0")
            database.execute("SELECT commission_type_snapshot,commission_rate_snapshot FROM trades LIMIT 0")
            tables = database.execute(
                """SELECT c.relname,c.relrowsecurity,
                    (has_table_privilege('anon',c.oid,'SELECT,INSERT,UPDATE,DELETE')
                     OR has_table_privilege('authenticated',c.oid,'SELECT,INSERT,UPDATE,DELETE')) AS browser_access
                FROM pg_class c JOIN pg_namespace n ON n.oid=c.relnamespace
                WHERE n.nspname='public' AND c.relkind='r'
                  AND c.relname IN ('settings','clients','trades','payments','uploads','upload_assets')"""
            ).fetchall()
            if len(tables) != 6 or any(not table["relrowsecurity"] or table["browser_access"] for table in tables):
                raise click.ClickException("The database schema or browser-access restrictions are incomplete. Run the SQL migration.")
            get_storage().check_bucket()
        except PsycopgError as exc:
            raise click.ClickException("PostgreSQL could not be verified. Check DATABASE_URL and apply the SQL migration.") from exc
        except (ValidationError, StorageUnavailable) as exc:
            raise click.ClickException(str(exc)) from exc
        click.echo("PostgreSQL is reachable, all six ledger tables have RLS and deny browser access, and Storage is private.")

    @app.cli.command("migrate-local")
    @click.option("--database", "database_path", required=True, type=click.Path(exists=True, dir_okay=False, path_type=Path))
    @click.option("--uploads", "upload_directory", required=True, type=click.Path(file_okay=False, path_type=Path))
    @click.option("--dry-run", is_flag=True, help="Validate the local source without connecting to cloud services.")
    def migrate_local_command(database_path, upload_directory, dry_run):
        from .transfer import inspect_local, migrate_local
        try:
            snapshot = inspect_local(database_path, upload_directory)
            if not dry_run:
                migrate_local(snapshot)
        except (ValidationError, StorageUnavailable, sqlite3.Error, OSError) as exc:
            raise click.ClickException(str(exc)) from exc
        except PsycopgError as exc:
            raise click.ClickException(
                "Cloud migration did not confirm completion. Check the cloud ledger before retrying; "
                "the local source is unchanged."
            ) from exc
        counts = ", ".join(f"{len(snapshot[table])} {table}" for table in ("clients", "trades", "payments", "uploads", "assets"))
        click.echo(("Preflight passed: " if dry_run else "Migration completed: ") + counts + ".")
        click.echo("The local database and original evidence were not modified. Cloud credentials and share links stay separate.")

    @app.cli.command("cleanup-uploads")
    @click.option("--dry-run", is_flag=True, help="Count abandoned uploads without deleting files or records.")
    def cleanup_uploads(dry_run):
        database = get_db()
        cutoff = int(time.time()) - 24 * 60 * 60
        where = "upload_id IS NULL AND expires_at<?"
        if dry_run:
            count = database.execute(f"SELECT COUNT(*) AS total FROM upload_assets WHERE {where}", (cutoff,)).fetchone()["total"]
            click.echo(f"{count} unsubmitted upload(s) expired more than 24 hours ago.")
            return
        removed = 0
        while True:
            locking = " FOR UPDATE SKIP LOCKED" if database.postgres else ""
            rows = database.execute(
                f"SELECT path FROM upload_assets WHERE {where} ORDER BY path LIMIT 100" + locking, (cutoff,)
            ).fetchall()
            if not rows:
                break
            paths = [row["path"] for row in rows]
            try:
                get_storage().remove(paths)
            except (StorageUnavailable, ValidationError, OSError) as exc:
                raise click.ClickException(str(exc)) from exc
            database.execute(
                "DELETE FROM upload_assets WHERE upload_id IS NULL AND path IN "
                f"({','.join('?' for _ in paths)})",
                paths,
            )
            database.commit()
            removed += len(paths)
        click.echo(f"Removed {removed} abandoned upload(s). Submitted evidence was not removed.")

    @app.route("/portal")
    @client_login_required
    def client_portal():
        cid = session["client_id"]
        client = get_client(cid)
        trades = query_trades({"client_id": cid})
        shares = get_client_calculated_holdings(cid, get_db(), sync_trades=False)

        prices = get_cached_prices(get_db())
        enriched_shares = []
        total_holdings_invested = ZERO
        total_holdings_current_value = ZERO

        for share in shares:
            s_dict = dict(share)
            qty = dec(s_dict.get("quantity"))
            buy_p = dec(s_dict.get("buy_price"))
            cost = qty * buy_p
            total_holdings_invested += cost

            stock_key = (s_dict.get("stock") or "").strip().upper()
            quote = prices.get(stock_key)
            if quote:
                cmp = quote["current_price"]
                cur_val = qty * cmp
                unrealized_pnl = cur_val - cost
                unrealized_pct = (unrealized_pnl / cost * 100) if cost else ZERO
                s_dict.update(
                    cmp=cmp,
                    current_value=cur_val,
                    unrealized_pnl=unrealized_pnl,
                    unrealized_pct=unrealized_pct,
                    price_change=quote.get("change", ZERO),
                    price_change_pct=quote.get("change_pct", ZERO),
                    price_updated_at=quote.get("updated_at"),
                )
                total_holdings_current_value += cur_val
            else:
                s_dict.update(
                    cmp=None,
                    current_value=cost,
                    unrealized_pnl=ZERO,
                    unrealized_pct=ZERO,
                    price_change=ZERO,
                    price_change_pct=ZERO,
                    price_updated_at=None,
                )
                total_holdings_current_value += cost
            enriched_shares.append(s_dict)

        total_unrealized_pnl = total_holdings_current_value - total_holdings_invested
        total_unrealized_pct = (total_unrealized_pnl / total_holdings_invested * 100) if total_holdings_invested else ZERO

        summary = summarize(trades)
        try:
            funds = get_db().execute("SELECT * FROM funds WHERE client_id=? ORDER BY added_on DESC, id DESC", (cid,)).fetchall()
        except Exception:
            funds = []
        total_funds = sum((dec(f["amount"]) for f in funds), ZERO)
        money_in_bank = total_funds - total_holdings_invested

        raw_undismissed_alerts = get_db().execute(
            """
            SELECT a.* FROM client_alerts a
            LEFT JOIN client_alert_views v ON v.alert_id=a.id AND v.client_id=?
            WHERE (a.client_id=? OR a.client_id IS NULL)
              AND a.active=1
              AND (v.acknowledged IS NULL OR v.acknowledged=0)
              AND (v.dismissed IS NULL OR v.dismissed=0)
            ORDER BY a.id ASC
            """,
            (cid, cid),
        ).fetchall()

        alias_map = get_alias_mapping(get_db())
        # For broadcast SELL recommendations, only deliver pop-up to clients who currently hold the stock
        held_canonical_stocks = {
            alias_map.get((sh.get("stock") or "").strip().upper(), (sh.get("stock") or "").strip().upper())
            for sh in shares
            if dec(sh.get("quantity") or 0) > ZERO
        }
        undismissed_alerts = []
        for a in raw_undismissed_alerts:
            if a["action"] == "SELL" and a["client_id"] is None:
                a_stock = alias_map.get((a.get("stock") or "").strip().upper(), (a.get("stock") or "").strip().upper())
                if a_stock not in held_canonical_stocks:
                    continue
            undismissed_alerts.append(a)

        return render_template(
            "portal.html",
            client=client,
            shares=enriched_shares,
            trades=trades,
            s=summary,
            funds=funds,
            total_funds=total_funds,
            money_in_bank=money_in_bank,
            total_holdings_invested=total_holdings_invested,
            total_holdings_current_value=total_holdings_current_value,
            total_unrealized_pnl=total_unrealized_pnl,
            total_unrealized_pct=total_unrealized_pct,
            undismissed_alerts=undismissed_alerts,
            public_token=portal_token(),
            today=date.today().isoformat(),
            known_stocks=get_known_stocks(get_db()),
        )

    @app.route("/portal/shares/new", methods=["POST"])
    @client_login_required
    def portal_share_add():
        cid = session["client_id"]
        db = get_db()
        stock = resolve_canonical_stock(db, request.form.get("stock", ""))
        if not stock:
            flash("Stock symbol is required.", "error")
            return redirect(url_for("client_portal"))
        try:
            quantity = int(round(number(request.form.get("quantity"), "Quantity", positive=True))) if request.form.get("quantity") else 1
            buy_price = number(request.form.get("buy_price"), "Buy Price") if request.form.get("buy_price") else None
            insert_trades(cid, [{
                "stock": stock,
                "quantity": quantity,
                "buy_price": buy_price,
                "sell_price": None,
                "buy_date": date.today().isoformat() if buy_price is not None else None,
                "notes": "Added from Client Portal",
            }], source="portal")

            # Update client_shares with weighted average price
            existing = db.execute(
                "SELECT id, quantity, buy_price FROM client_shares WHERE client_id=? AND UPPER(TRIM(stock))=?",
                (cid, stock.upper())
            ).fetchone()
            if existing:
                old_qty = dec(existing["quantity"] or 0)
                old_bp = dec(existing["buy_price"] or 0)
                add_qty = dec(quantity)
                add_bp = dec(buy_price or old_bp)
                new_qty = int(round(old_qty + add_qty))
                if new_qty > 0 and (old_qty * old_bp + add_qty * add_bp) > 0:
                    new_avg = ((old_qty * old_bp + add_qty * add_bp) / dec(new_qty)).quantize(CENT, rounding=ROUND_HALF_UP)
                else:
                    new_avg = add_bp
                db.execute("UPDATE client_shares SET quantity=?, buy_price=? WHERE id=?", (new_qty, new_avg, existing["id"]))
            else:
                db.execute(
                    "INSERT INTO client_shares (client_id, stock, quantity, buy_price, added_on) VALUES (?, ?, ?, ?, ?)",
                    (cid, stock, quantity, buy_price, date.today().isoformat())
                )

            flash(f"Trade for {stock} recorded. Your holdings have been updated.", "success")
        except Exception as exc:
            current_app.logger.warning("Failed to add portal share: %s", exc)
            flash(f"Could not record holding: {exc}", "error")
        return redirect(url_for("client_portal"))

    @app.route("/portal/trades/new", methods=["POST"])
    @client_login_required
    def portal_trade_add():
        cid = session["client_id"]
        client = get_client(cid)
        db = get_db()
        raw_stock = request.form.get("stock", "").strip()
        stock = resolve_canonical_stock(db, raw_stock) if raw_stock else ""
        if not stock:
            flash("Stock symbol is required.", "error")
            return redirect(url_for("client_portal"))

        try:
            quantity = int(round(number(request.form.get("quantity"), "Quantity", required=True, positive=True)))
        except Exception as exc:
            flash(f"Invalid quantity: {exc}", "error")
            return redirect(url_for("client_portal"))

        raw_bp = (request.form.get("buy_price") or "").strip()
        raw_sp = (request.form.get("sell_price") or "").strip()

        buy_price = None
        if raw_bp:
            try:
                buy_price = number(raw_bp, "Buy Price")
            except Exception as exc:
                flash(f"Invalid buy price: {exc}", "error")
                return redirect(url_for("client_portal"))

        sell_price = None
        if raw_sp:
            try:
                sell_price = number(raw_sp, "Sell Price")
            except Exception as exc:
                flash(f"Invalid sell price: {exc}", "error")
                return redirect(url_for("client_portal"))

        # Both buy and sell cannot be empty
        if buy_price is None and sell_price is None:
            flash("Please enter either a Buy Price or a Sell Price. Both cannot be empty.", "error")
            return redirect(url_for("client_portal"))

        raw_bdate = (request.form.get("buy_date") or "").strip()
        buy_date = valid_date(raw_bdate, "Buy Date") if raw_bdate else (date.today().isoformat() if buy_price is not None else None)

        raw_sdate = (request.form.get("sell_date") or "").strip()
        sell_date = valid_date(raw_sdate, "Sell Date") if raw_sdate else (date.today().isoformat() if sell_price is not None else None)

        user_notes = text(request.form.get("notes"), "Notes", limit=200) or ""

        # Handle optional proof image upload
        proof_file = request.files.get("proof_image")
        upload_id = None
        if proof_file and proof_file.filename:
            content = proof_file.read()
            if len(content) > 0:
                mimetype = proof_file.mimetype or "image/png"
                if mimetype not in IMAGE_TYPES:
                    ext = os.path.splitext(proof_file.filename)[1].lower()
                    if ext in (".jpg", ".jpeg"):
                        mimetype = "image/jpeg"
                    elif ext == ".png":
                        mimetype = "image/png"
                    elif ext == ".webp":
                        mimetype = "image/webp"
                    else:
                        mimetype = "image/png"

                try:
                    validate_image(content, mimetype)
                    storage_path = "evidence/" + uuid.uuid4().hex + IMAGE_TYPES[mimetype][1]
                    storage = get_storage()
                    storage.store(storage_path, content, mimetype)

                    sub_uuid = str(uuid.uuid4())
                    sub_hash = hashlib.sha256(content).hexdigest()
                    owner = f"client:{cid}"
                    up_row = db.execute(
                        """INSERT INTO uploads (
                            client_name, client_id, filename, platform, trade_count,
                            submission_id, submission_hash, submission_owner
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?) RETURNING id""",
                        (client["name"], cid, storage_path, "Client Portal Add Trade", 1, sub_uuid, sub_hash, owner),
                    ).fetchone()
                    upload_id = up_row["id"]

                    db.execute(
                        """INSERT INTO upload_assets (path, owner_hash, scope_hash, content_type, byte_size, expires_at, upload_id)
                        VALUES (?, ?, ?, ?, ?, ?, ?)""",
                        (storage_path, owner, "portal", mimetype, len(content), int(time.time()) + 365 * 86400, upload_id),
                    )
                except Exception as exc:
                    current_app.logger.warning("Trade screenshot upload skipped: %s", exc)

        # -------------------------------------------------------------
        # HOLDINGS & TRADE LOGIC:
        # 1. When selling and guest hasn't entered buy price:
        #    Infer buy price from the client's current holding average price,
        #    so realized P&L is calculated as sell_price - avg_price.
        #    Deduct sold quantity from holdings.
        # 2. When client adds a new buy trade:
        #    Update the holding in client_shares, recalculating weighted avg price.
        # -------------------------------------------------------------
        existing_holding = db.execute(
            "SELECT id, quantity, buy_price, added_on FROM client_shares WHERE client_id=? AND UPPER(TRIM(stock))=?",
            (cid, stock.upper())
        ).fetchone()

        if sell_price is not None and buy_price is None:
            # Client sold shares without entering buy price
            if existing_holding and existing_holding["buy_price"] is not None:
                holding_avg_bp = dec(existing_holding["buy_price"])
                buy_price = holding_avg_bp
                buy_date = (existing_holding["added_on"] or "")[:10] or date.today().isoformat()
                user_notes = user_notes or "Sold from Holdings"

                old_h_qty = dec(existing_holding["quantity"] or 0)
                rem_qty = int(round(old_h_qty - dec(quantity)))
                if rem_qty <= 0:
                    db.execute("DELETE FROM client_shares WHERE id=?", (existing_holding["id"],))
                else:
                    db.execute("UPDATE client_shares SET quantity=? WHERE id=?", (rem_qty, existing_holding["id"]))
            else:
                # Check if there is an open trade in trades for this stock
                open_t = db.execute(
                    """SELECT id, quantity, buy_price, buy_date FROM trades
                       WHERE client_id=? AND UPPER(TRIM(stock))=? AND buy_price IS NOT NULL AND (sell_price IS NULL OR sell_price=0)
                       ORDER BY id ASC LIMIT 1""",
                    (cid, stock.upper())
                ).fetchone()
                if open_t and open_t["buy_price"] is not None:
                    buy_price = dec(open_t["buy_price"])
                    buy_date = open_t["buy_date"] or date.today().isoformat()
                    ot_qty = dec(open_t["quantity"] or 0)
                    rem_ot_qty = int(round(ot_qty - dec(quantity)))
                    if rem_ot_qty <= 0:
                        db.execute("DELETE FROM trades WHERE id=?", (open_t["id"],))
                    else:
                        db.execute("UPDATE trades SET quantity=? WHERE id=?", (rem_ot_qty, open_t["id"]))

        elif sell_price is not None and buy_price is not None:
            # Completed trade with both prices entered
            if existing_holding:
                old_h_qty = dec(existing_holding["quantity"] or 0)
                rem_qty = int(round(old_h_qty - dec(quantity)))
                if rem_qty <= 0:
                    db.execute("DELETE FROM client_shares WHERE id=?", (existing_holding["id"],))
                else:
                    db.execute("UPDATE client_shares SET quantity=? WHERE id=?", (rem_qty, existing_holding["id"]))

        elif buy_price is not None and sell_price is None:
            # Buy-only trade: update holdings with new weighted average price
            if existing_holding:
                old_qty = dec(existing_holding["quantity"] or 0)
                old_bp = dec(existing_holding["buy_price"] or 0)
                add_qty = dec(quantity)
                add_bp = dec(buy_price)
                new_qty = int(round(old_qty + add_qty))
                if new_qty > 0 and (old_qty * old_bp + add_qty * add_bp) > 0:
                    new_avg = ((old_qty * old_bp + add_qty * add_bp) / dec(new_qty)).quantize(CENT, rounding=ROUND_HALF_UP)
                else:
                    new_avg = add_bp
                db.execute(
                    "UPDATE client_shares SET quantity=?, buy_price=? WHERE id=?",
                    (new_qty, new_avg, existing_holding["id"])
                )
            else:
                db.execute(
                    "INSERT INTO client_shares (client_id, stock, quantity, buy_price, added_on) VALUES (?, ?, ?, ?, ?)",
                    (cid, stock, quantity, buy_price, buy_date or date.today().isoformat())
                )

        # Pair constraints for DB:
        if buy_price is not None and not buy_date:
            buy_date = date.today().isoformat()
        if sell_price is not None and not sell_date:
            sell_date = date.today().isoformat()

        # Satisfy DB date order constraints
        if buy_date and sell_date and sell_date < buy_date:
            sell_date = buy_date

        try:
            trade_ids = insert_trades(cid, [{
                "stock": stock,
                "quantity": quantity,
                "buy_price": buy_price,
                "sell_price": sell_price,
                "buy_date": buy_date,
                "sell_date": sell_date,
                "notes": user_notes or "Added via Client Portal",
            }], source="portal", upload_id=upload_id)
            trade_id = trade_ids[0] if trade_ids else None
        except Exception as exc:
            current_app.logger.exception("Failed to insert portal trade: %s", exc)
            flash(f"Could not record trade: {exc}", "error")
            return redirect(url_for("client_portal"))

        # Eagerly fetch live price quote for the stock
        try:
            quote = fetch_stock_quote(stock)
            if quote:
                db.execute(
                    """
                    INSERT INTO stock_prices (stock, current_price, previous_close, change, change_pct, updated_at)
                    VALUES (?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
                    ON CONFLICT(stock) DO UPDATE SET
                        current_price=excluded.current_price,
                        previous_close=excluded.previous_close,
                        change=excluded.change,
                        change_pct=excluded.change_pct,
                        updated_at=CURRENT_TIMESTAMP
                    """,
                    (stock, quote["current_price"], quote.get("previous_close"), quote.get("change"), quote.get("change_pct")),
                )
        except Exception:
            pass

        # If trade is completed (has both buy_price and sell_price), add realised profit to bank funds ONCE
        if buy_price is not None and sell_price is not None and trade_id:
            pnl = ((dec(sell_price) - dec(buy_price)) * dec(quantity)).quantize(CENT, rounding=ROUND_HALF_UP)
            if pnl != ZERO:
                tag = f"[Trade #{trade_id}]"
                existing = db.execute(
                    "SELECT id FROM funds WHERE client_id=? AND notes LIKE ?",
                    (cid, f"%{tag}%")
                ).fetchone()
                if not existing:
                    s_date = sell_date or date.today().isoformat()
                    pnl_note = f"[Realized {'Profit' if pnl > 0 else 'Loss'}] {stock} ({quantity} sh) {tag}"
                    db.execute(
                        "INSERT INTO funds (client_id, amount, added_on, notes) VALUES (?, ?, ?, ?)",
                        (cid, pnl, s_date, pnl_note),
                    )
            if pnl >= ZERO:
                flash(f"Trade completed for {stock}. Realized profit of ₹{pnl:,.2f} added to your bank funds.", "success")
            else:
                flash(f"Trade completed for {stock}. Realized loss of ₹{abs(pnl):,.2f} recorded in your funds ledger.", "success")
        elif buy_price is not None:
            flash(f"Buy trade for {quantity} shares of {stock} recorded. Your holdings and average buy price have been updated.", "success")
        else:
            flash(f"Sell trade for {quantity} shares of {stock} recorded. Your holdings have been updated.", "success")

        invalidate_prices_cache()
        invalidate_stock_caches()
        db.commit()
        return redirect(url_for("client_portal"))

    @app.route("/portal/trades/<int:trade_id>/delete-screenshot", methods=["POST"])
    @client_login_required
    def portal_trade_delete_screenshot(trade_id):
        cid = session["client_id"]
        db = get_db()
        trade = db.execute("SELECT * FROM trades WHERE id=? AND client_id=?", (trade_id, cid)).fetchone()
        if not trade:
            flash("Trade not found or unauthorized.", "error")
            return redirect(url_for("client_portal"))

        upload_id = trade["upload_id"]
        if not upload_id:
            flash("No screenshot attached to this trade.", "info")
            return redirect(url_for("client_portal"))

        stock_name = trade["stock"]
        db.execute("UPDATE trades SET upload_id=NULL WHERE id=? AND client_id=?", (trade_id, cid))

        other_trades = db.execute("SELECT COUNT(*) AS c FROM trades WHERE upload_id=?", (upload_id,)).fetchone()
        count = other_trades["c"] if other_trades else 0
        if count == 0:
            try:
                assets = db.execute("SELECT path FROM upload_assets WHERE upload_id=?", (upload_id,)).fetchall()
                paths = [a["path"] for a in assets if a and a["path"]]
                if paths:
                    try:
                        storage = get_storage()
                        storage.remove(paths)
                    except Exception as storage_exc:
                        current_app.logger.warning("Failed to remove storage files for upload %s: %s", upload_id, storage_exc)
                db.execute("DELETE FROM upload_assets WHERE upload_id=?", (upload_id,))
                db.execute("DELETE FROM uploads WHERE id=?", (upload_id,))
            except Exception as exc:
                current_app.logger.warning("Error cleaning up orphaned upload record %s: %s", upload_id, exc)

        db.commit()
        flash(f"Proof screenshot removed from {stock_name}.", "success")
        return redirect(url_for("client_portal"))

    @app.route("/portal/funds/adjust", methods=["POST"])
    @client_login_required
    def portal_funds_adjust():
        cid = session["client_id"]
        client = get_client(cid)
        action = request.form.get("action", "add").strip().lower()
        amount = number(request.form.get("amount"), "Amount", required=True, positive=True, places=2)
        added_on = valid_date(request.form.get("date"), "Date") or date.today().isoformat()
        user_notes = text(request.form.get("notes"), "Notes", limit=200) or ""

        if action == "remove":
            fund_amount = -amount
            note_prefix = "[Client Withdrawal]"
            action_text = f"Withdrew ₹{amount:,.2f}"
        else:
            fund_amount = amount
            note_prefix = "[Client Deposit]"
            action_text = f"Added ₹{amount:,.2f}"

        notes_entry = f"{note_prefix} {user_notes}".strip()

        get_db().execute(
            "INSERT INTO funds (client_id, amount, added_on, notes) VALUES (?, ?, ?, ?)",
            (cid, fund_amount, added_on, notes_entry),
        )
        flash(f"{action_text} recorded in your funds ledger.", "success")
        return redirect(url_for("client_portal"))

    @app.route("/portal/set-username", methods=["POST"])
    @client_login_required
    def portal_set_username():
        cid = session["client_id"]
        client = get_client(cid)
        # Client can only set their username one time only
        if client.get("username"):
            flash("Your username has already been set and is permanently locked. Only your administrator can modify it.", "error")
            return redirect(url_for("client_portal"))

        raw_username = request.form.get("username", "").strip()
        username = raw_username.lstrip("@").strip().lower()
        if not username:
            flash("Please enter a username.", "error")
            return redirect(url_for("client_portal"))

        if len(username) < 3 or len(username) > 30:
            flash("Username must be between 3 and 30 characters.", "error")
            return redirect(url_for("client_portal"))

        if not re.match(r"^[a-zA-Z0-9_.]+$", username):
            flash("Username can only contain letters, numbers, underscores, and dots.", "error")
            return redirect(url_for("client_portal"))

        existing = get_db().execute("SELECT id FROM clients WHERE lower(username)=lower(?) AND id<>?", (username, cid)).fetchone()
        if existing:
            flash(f"The username @{username} is already taken. Please choose another one.", "error")
            return redirect(url_for("client_portal"))

        get_db().execute("UPDATE clients SET username=? WHERE id=?", (username, cid))
        flash(f"Your username @{username} has been saved and locked! You can now use this handle to sign into your portfolio.", "success")
        return redirect(url_for("client_portal"))

    @app.route("/portal/password/update", methods=["POST"])
    @client_login_required
    def client_update_password():
        cid = session["client_id"]
        new_pwd = request.form.get("new_password", "").strip()
        confirm_pwd = request.form.get("confirm_password", "").strip()

        if len(new_pwd) < 6:
            flash("New password must be at least 6 characters long.", "error")
            return redirect(url_for("client_portal"))

        if new_pwd != confirm_pwd:
            flash("Passwords do not match. Please re-enter carefully.", "error")
            return redirect(url_for("client_portal"))

        p_hash = generate_password_hash(new_pwd)
        get_db().execute(
            "UPDATE clients SET password_hash=?, password_plain=?, password_changed=1 WHERE id=?",
            (p_hash, new_pwd, cid),
        )
        flash("Password updated successfully! Keep this password safe for your future logins.", "success")
        return redirect(url_for("client_portal"))

    @app.route("/api/alerts/<int:aid>/acknowledge", methods=["POST"])
    @app.route("/api/alerts/<int:aid>/dismiss", methods=["POST"])
    @client_login_required
    def api_alert_acknowledge(aid):
        cid = session["client_id"]
        get_db().execute(
            """
            INSERT INTO client_alert_views (alert_id, client_id, dismissed, acknowledged, acknowledged_at, viewed_at)
            VALUES (?, ?, 1, 1, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
            ON CONFLICT(alert_id, client_id) DO UPDATE SET 
                dismissed=1,
                acknowledged=1,
                acknowledged_at=CURRENT_TIMESTAMP,
                viewed_at=CURRENT_TIMESTAMP
            """,
            (aid, cid),
        )
        return jsonify(success=True, acknowledged=True)

    @app.route("/api/stocks/suggestions", methods=["GET"])
    def api_stock_suggestions():
        return jsonify(get_stock_suggestions(get_db()))

    @app.route("/alerts", methods=["GET"])
    @login_required
    def alerts_list():
        db = get_db()
        alerts = db.execute(
            """
            SELECT a.*, c.name as client_name,
                   (SELECT COUNT(*) FROM client_alert_views v WHERE v.alert_id=a.id AND (v.dismissed=1 OR v.acknowledged=1)) as view_count
            FROM client_alerts a
            LEFT JOIN clients c ON c.id=a.client_id
            ORDER BY a.created_at DESC
            """
        ).fetchall()
        clients = db.execute("SELECT id, name FROM clients WHERE active=1 ORDER BY lower(name)").fetchall()

        views = db.execute(
            """
            SELECT v.alert_id, v.client_id, v.acknowledged_at, v.viewed_at, c.name as client_name
            FROM client_alert_views v
            JOIN clients c ON c.id = v.client_id
            WHERE (v.dismissed = 1 OR v.acknowledged = 1)
            ORDER BY lower(c.name)
            """
        ).fetchall()

        # Batch-fetch all active clients' non-zero holdings in ONE query to avoid N+1 queries per SELL alert
        all_holdings = db.execute(
            """
            SELECT h.client_id, c.name as client_name, h.stock
            FROM client_shares h
            JOIN clients c ON c.id = h.client_id
            WHERE c.active = 1 AND h.quantity > 0
            ORDER BY lower(c.name)
            """
        ).fetchall()

        alias_map = get_alias_mapping(db)
        stock_holders_map = {}
        for h in all_holdings:
            sym = (h["stock"] or "").strip().upper()
            canon = alias_map.get(sym, sym)
            stock_holders_map.setdefault(canon, {})[h["client_id"]] = h["client_name"]

        ack_by_alert = {}
        for v in views:
            ack_by_alert.setdefault(v["alert_id"], []).append(v)

        active_clients_list = [dict(c) for c in clients]
        active_client_map = {c["id"]: c["name"] for c in clients}

        enriched_alerts = []
        for a in alerts:
            ad = dict(a)
            acks = ack_by_alert.get(ad["id"], [])
            acked_client_ids = {v["client_id"] for v in acks}
            acked_clients = [
                {
                    "id": v["client_id"],
                    "name": v["client_name"],
                    "acknowledged_at": v.get("acknowledged_at") or v.get("viewed_at"),
                }
                for v in acks
            ]

            if ad.get("client_id"):
                cid = ad["client_id"]
                cname = ad.get("client_name") or active_client_map.get(cid, f"Client #{cid}")
                target_clients = [{"id": cid, "name": cname}]
                if cid in acked_client_ids:
                    pending_clients = []
                else:
                    pending_clients = [{"id": cid, "name": cname}]
            elif ad.get("action") == "SELL":
                # Sell pop-up broadcast: only sent to clients who hold the given stock
                canon_stk = alias_map.get((ad.get("stock") or "").strip().upper(), (ad.get("stock") or "").strip().upper())
                holder_map = dict(stock_holders_map.get(canon_stk, {}))
                for v in acks:
                    if v["client_id"] not in holder_map:
                        holder_map[v["client_id"]] = v["client_name"]
                target_clients = [{"id": cid, "name": holder_map[cid]} for cid in sorted(holder_map.keys(), key=lambda i: holder_map[i].lower())]
                pending_clients = [c for c in target_clients if c["id"] not in acked_client_ids]
            else:
                target_clients = active_clients_list
                pending_clients = [c for c in active_clients_list if c["id"] not in acked_client_ids]

            ad["target_count"] = len(target_clients)
            ad["ack_count"] = len(acked_clients)
            ad["acknowledged_clients"] = acked_clients
            ad["pending_clients"] = pending_clients
            ad["all_acknowledged"] = (len(target_clients) > 0 and len(pending_clients) == 0)
            ad["is_sell_broadcast"] = (ad.get("action") == "SELL" and not ad.get("client_id"))
            enriched_alerts.append(ad)

        ensure_push_tables(db)
        try:
            push_subs_count = db.execute("SELECT COUNT(*) AS c FROM client_push_subscriptions").fetchone()["c"]
        except Exception:
            push_subs_count = 0

        return render_template("alerts.html", alerts=enriched_alerts, clients=clients, push_subs_count=push_subs_count)

    def ensure_push_tables(db):
        try:
            if db.postgres:
                db.execute("""
                CREATE TABLE IF NOT EXISTS client_push_subscriptions (
                    id BIGSERIAL PRIMARY KEY,
                    client_id BIGINT REFERENCES clients(id) ON DELETE CASCADE,
                    endpoint TEXT NOT NULL UNIQUE,
                    p256dh TEXT NOT NULL,
                    auth TEXT NOT NULL,
                    user_agent TEXT,
                    created_at TIMESTAMP WITH TIME ZONE DEFAULT CURRENT_TIMESTAMP
                )
                """)
                db.execute("CREATE INDEX IF NOT EXISTS idx_push_client ON client_push_subscriptions(client_id)")
            else:
                db.execute("""
                CREATE TABLE IF NOT EXISTS client_push_subscriptions (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    client_id INTEGER REFERENCES clients(id) ON DELETE CASCADE,
                    endpoint TEXT NOT NULL UNIQUE,
                    p256dh TEXT NOT NULL,
                    auth TEXT NOT NULL,
                    user_agent TEXT,
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                )
                """)
                db.execute("CREATE INDEX IF NOT EXISTS idx_push_client ON client_push_subscriptions(client_id)")
        except Exception as exc:
            current_app.logger.warning("ensure_push_tables notice: %s", exc)

    def dispatch_web_push(client_id=None, title="Trade Alert", body="New trade recommendation", url="/portal", stock=None, action=None):
        db = get_db()
        ensure_push_tables(db)
        try:
            if client_id is not None:
                subs = db.execute("SELECT id, endpoint, p256dh, auth FROM client_push_subscriptions WHERE client_id=?", (client_id,)).fetchall()
            else:
                subs = db.execute("SELECT id, endpoint, p256dh, auth FROM client_push_subscriptions").fetchall()
        except Exception as exc:
            current_app.logger.warning("Push query failed: %s", exc)
            return 0

        if not subs:
            return 0

        try:
            from pywebpush import webpush, WebPushException
        except ImportError:
            current_app.logger.warning("pywebpush not installed, skipping push dispatch")
            return 0

        vapid_priv = app.config.get("VAPID_PRIVATE_KEY")
        vapid_claims = {"sub": app.config.get("VAPID_CLAIM_EMAIL", "mailto:admin@commissionbook.com")}

        payload = json.dumps({
            "title": title,
            "body": body,
            "icon": "/static/icon-192.png",
            "badge": "/static/icon-192.png",
            "url": url,
            "tag": f"trade-{stock or 'alert'}",
            "data": {"url": url}
        })

        sent_count = 0
        expired_ids = []
        for sub in subs:
            sub_info = {
                "endpoint": sub["endpoint"],
                "keys": {
                    "p256dh": sub["p256dh"],
                    "auth": sub["auth"],
                }
            }
            try:
                webpush(
                    subscription_info=sub_info,
                    data=payload,
                    vapid_private_key=vapid_priv,
                    vapid_claims=vapid_claims,
                    ttl=86400
                )
                sent_count += 1
            except WebPushException as ex:
                current_app.logger.warning("WebPush send failed: %s", ex)
                if hasattr(ex, "response") and ex.response is not None and ex.response.status_code in (404, 410):
                    expired_ids.append(sub["id"])
            except Exception as ex:
                current_app.logger.warning("WebPush unexpected error: %s", ex)

        if expired_ids:
            try:
                for eid in expired_ids:
                    db.execute("DELETE FROM client_push_subscriptions WHERE id=?", (eid,))
                db.commit()
            except Exception:
                pass

        return sent_count

    @app.route("/sw.js")
    def service_worker():
        response = send_from_directory(app.static_folder, "sw.js", mimetype="application/javascript")
        response.headers["Service-Worker-Allowed"] = "/"
        response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
        return response

    @app.route("/api/push/public-key")
    def push_public_key():
        return jsonify(publicKey=app.config.get("VAPID_PUBLIC_KEY", ""))

    @app.route("/portal/push/subscribe", methods=["POST"])
    @client_login_required
    def portal_push_subscribe():
        cid = session["client_id"]
        data = request.get_json(silent=True) or {}
        endpoint = (data.get("endpoint") or "").strip()
        keys = data.get("keys") or {}
        p256dh = (keys.get("p256dh") or "").strip()
        auth = (keys.get("auth") or "").strip()
        user_agent = (request.headers.get("User-Agent") or "")[:250]

        if not endpoint or not p256dh or not auth:
            return jsonify(ok=False, error="Invalid subscription payload"), 400

        db = get_db()
        ensure_push_tables(db)
        if db.postgres:
            db.execute(
                """
                INSERT INTO client_push_subscriptions (client_id, endpoint, p256dh, auth, user_agent)
                VALUES (?, ?, ?, ?, ?)
                ON CONFLICT (endpoint) DO UPDATE SET
                    client_id = EXCLUDED.client_id,
                    p256dh = EXCLUDED.p256dh,
                    auth = EXCLUDED.auth,
                    user_agent = EXCLUDED.user_agent,
                    created_at = CURRENT_TIMESTAMP
                """,
                (cid, endpoint, p256dh, auth, user_agent)
            )
        else:
            db.execute(
                """
                INSERT INTO client_push_subscriptions (client_id, endpoint, p256dh, auth, user_agent)
                VALUES (?, ?, ?, ?, ?)
                ON CONFLICT (endpoint) DO UPDATE SET
                    client_id = excluded.client_id,
                    p256dh = excluded.p256dh,
                    auth = excluded.auth,
                    user_agent = excluded.user_agent,
                    created_at = CURRENT_TIMESTAMP
                """,
                (cid, endpoint, p256dh, auth, user_agent)
            )
        db.commit()
        return jsonify(ok=True, message="Push notifications enabled")

    @app.route("/portal/push/unsubscribe", methods=["POST"])
    @client_login_required
    def portal_push_unsubscribe():
        data = request.get_json(silent=True) or {}
        endpoint = (data.get("endpoint") or "").strip()
        if endpoint:
            db = get_db()
            ensure_push_tables(db)
            db.execute("DELETE FROM client_push_subscriptions WHERE endpoint=?", (endpoint,))
            db.commit()
        return jsonify(ok=True)

    @app.route("/admin/push/test", methods=["POST"])
    @login_required
    def admin_push_test():
        client_id_raw = (request.form.get("client_id") or "").strip()
        client_id = int(client_id_raw) if client_id_raw and client_id_raw != "all" else None
        target_name = get_client(client_id)["name"] if client_id else "All Subscribed Clients"
        sent = dispatch_web_push(
            client_id=client_id,
            title="🔔 Commission Book Test Alert",
            body="Push notifications are active and working on this device!",
            url="/portal"
        )
        if sent > 0:
            flash(f"Test push notification dispatched to {sent} active device(s) for {target_name}.", "success")
        else:
            flash(f"No active push subscriptions found for {target_name}. Ask the client to tap 'Enable Alerts' on their portal first.", "info")
        return redirect(safe_redirect_target(request.referrer, url_for("alerts_list")))

    @app.route("/alerts/new", methods=["POST"])
    @login_required
    def alert_create():
        action = request.form.get("action", "BUY").upper()
        if action not in ("BUY", "SELL"):
            abort(400)
        stock = request.form.get("stock", "").strip().upper()
        back = request.form.get("back")
        if not stock:
            flash("Stock symbol is required.", "error")
            return redirect(safe_redirect_target(back, url_for("alerts_list")))
        price = number(request.form.get("price"), "Price", positive=True) if request.form.get("price") else None
        target_price = number(request.form.get("target_price"), "Target Price", positive=True) if request.form.get("target_price") else None
        stop_loss = number(request.form.get("stop_loss"), "Stop Loss", positive=True) if request.form.get("stop_loss") else None
        message = request.form.get("message", "").strip()
        client_id_raw = request.form.get("client_id", "").strip()
        client_id = int(client_id_raw) if client_id_raw and client_id_raw != "all" else None
        canonical_stock = resolve_canonical_stock(get_db(), stock)

        get_db().execute(
            """
            INSERT INTO client_alerts (client_id, action, stock, price, target_price, stop_loss, message, active)
            VALUES (?, ?, ?, ?, ?, ?, ?, 1)
            """,
            (client_id, action, canonical_stock, price, target_price, stop_loss, message),
        )

        send_push = request.form.get("send_push", "1") == "1"
        push_sent = 0
        if send_push:
            action_verb = "BUY" if action == "BUY" else "SELL"
            push_title = f"📢 Trade Advisory: {action_verb} {canonical_stock}"
            details = []
            if price:
                details.append(f"Price: ₹{price:,.2f}")
            if target_price:
                details.append(f"Target: ₹{target_price:,.2f}")
            if stop_loss:
                details.append(f"SL: ₹{stop_loss:,.2f}")
            push_body = " | ".join(details) if details else (message or "New trade recommendation posted.")
            if message and details:
                push_body += f"\n{message}"

            if client_id is not None:
                push_sent = dispatch_web_push(client_id=client_id, title=push_title, body=push_body, stock=canonical_stock, action=action)
            elif action == "SELL":
                holding_clients = get_clients_holding_stock(canonical_stock, get_db())
                for hc in holding_clients:
                    push_sent += dispatch_web_push(client_id=hc["id"], title=push_title, body=push_body, stock=canonical_stock, action=action)
            else:
                push_sent = dispatch_web_push(client_id=None, title=push_title, body=push_body, stock=canonical_stock, action=action)

        push_note = f" (Pushed to {push_sent} device{'s' if push_sent != 1 else ''})" if push_sent > 0 else ""
        if client_id is not None:
            recipient = get_client(client_id)["name"]
            flash(f"Pop-up trade for {canonical_stock} ({action}) sent to {recipient}{push_note}.", "success")
        elif action == "SELL":
            holding_clients = get_clients_holding_stock(canonical_stock, get_db())
            if holding_clients:
                names = ", ".join(h["name"] for h in holding_clients[:3])
                suffix = f" (+{len(holding_clients) - 3} more)" if len(holding_clients) > 3 else ""
                flash(f"Pop-up SELL advisory for {canonical_stock} sent to {len(holding_clients)} client(s) holding this stock: {names}{suffix}{push_note}.", "success")
            else:
                flash(f"Pop-up SELL advisory for {canonical_stock} recorded. Notice: No clients currently hold this stock, so no client portals will be prompted until holdings exist.", "info")
        else:
            recipient = "All clients (Broadcast)"
            flash(f"Pop-up trade for {canonical_stock} ({action}) sent to {recipient}{push_note}.", "success")
        return redirect(safe_redirect_target(back, url_for("alerts_list")))

    @app.route("/alerts/<int:aid>/toggle", methods=["POST"])
    @login_required
    def alert_toggle(aid):
        alert = get_db().execute("SELECT * FROM client_alerts WHERE id=?", (aid,)).fetchone()
        if not alert:
            abort(404)
        new_status = 0 if alert["active"] else 1
        get_db().execute("UPDATE client_alerts SET active=? WHERE id=?", (new_status, aid))
        flash("Alert status updated.", "success")
        return redirect(url_for("alerts_list"))

    @app.route("/alerts/<int:aid>/delete", methods=["POST"])
    @login_required
    def alert_delete(aid):
        get_db().execute("DELETE FROM client_alerts WHERE id=?", (aid,))
        flash("Alert deleted.", "success")
        return redirect(url_for("alerts_list"))

    if not (config and config.get("TESTING")):
        scheduler = DailyPriceScheduler(app)
        scheduler.start()
        app.extensions["price_scheduler"] = scheduler

    return app
