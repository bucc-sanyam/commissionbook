import io
import json
import os
import sqlite3
import tempfile
import threading
import time
import unittest
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from unittest.mock import MagicMock, patch
from urllib.parse import urlsplit

import httpx
import psycopg
from flask import Flask, template_rendered
from openpyxl import Workbook, load_workbook
from PIL import Image
from werkzeug.security import generate_password_hash

from backend.database import Database, Result, get_db
from backend.security import configure_app
from backend.storage import StorageUnavailable, SupabaseStorage
from backend.transfer import inspect_local, migrate_local
from backend.validation import ValidationError, number
from backend.web import create_app, get_or_create_client, query_trades, set_setting
from trade_parser import parse_text


PASSWORD = "temporary-password"
PASSWORD_HASH = generate_password_hash(PASSWORD, method="pbkdf2:sha256:1000")
PORTAL = "isolated-test-portal-token"
CSRF = "isolated-test-csrf-token"
ROW = {
    "stock": "INFY", "quantity": 2, "buy_price": 100, "sell_price": 150,
    "buy_date": "2026-01-02", "sell_date": "2026-01-05", "notes": "",
}


def png_bytes():
    output = io.BytesIO()
    Image.new("RGB", (4, 4), color="white").save(output, format="PNG")
    return output.getvalue()


class BackendTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="commission-backend-")
        self.root = Path(self.directory.name)
        self.config = {
            "TESTING": True, "PRODUCTION": False, "BACKEND_MODE": "local",
            "DATABASE": str(self.root / "ledger.db"), "DATABASE_URL": "",
            "UPLOAD_DIR": str(self.root / "uploads"), "SECRET_KEY": "test-secret-" * 5,
            "ADMIN_PASSWORD": "", "ADMIN_PASSWORD_HASH": PASSWORD_HASH,
            "SUPABASE_URL": "", "SUPABASE_SERVICE_ROLE_KEY": "",
            "SUPABASE_STORAGE_BUCKET": "commission-evidence",
            "PUBLIC_BASE_URL": "", "SESSION_COOKIE_SECURE": False,
        }
        self.environment = patch.dict(os.environ, {}, clear=True)
        self.environment.start()
        self.app = create_app(self.config)
        self.client = self.app.test_client()
        self.authorize_client(self.client)
        with self.app.app_context():
            set_setting("upload_token", PORTAL)
            get_db().commit()

    def tearDown(self):
        self.environment.stop()
        self.directory.cleanup()

    def authorize_client(self, client, admin=False):
        with client.session_transaction() as session:
            session["_csrf"] = CSRF
            if admin:
                session["admin"] = True
                session["auth_version"] = self.app.config["ADMIN_AUTH_VERSION"]

    def admin(self):
        self.authorize_client(self.client, admin=True)

    def post(self, path, data=None, client=None, **kwargs):
        return (client or self.client).post(path, data=data, headers={"X-CSRF-Token": CSRF}, **kwargs)

    def payload(self, **changes):
        return {
            "name": "Alice", "rows": [dict(ROW)], "portal_token": PORTAL, "code": "",
            "ocr_text": "", "platform": "generic", "images": [], "submission_id": str(uuid.uuid4()),
            **changes,
        }

    def submit(self, payload=None, client=None):
        return self.post("/api/submit", json=payload or self.payload(), client=client)

    def records(self, sql, parameters=()):
        with self.app.app_context():
            return get_db().execute(sql, parameters).fetchall()

    def sign(self, image=None):
        image = png_bytes() if image is None else image
        response = self.post("/api/uploads/sign", json={
            "portal_token": PORTAL, "code": "", "filename": "screenshot.png",
            "content_type": "image/png", "size": len(image),
        })
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        return response.json, image

    def upload(self, signed, contents):
        return self.client.put(
            urlsplit(signed["upload_url"]).path, data=contents, headers=signed["headers"],
        )

    def test_private_routes_and_legacy_upload_do_not_reveal_the_portal(self):
        for path in ("/", "/trades", "/clients", "/payments", "/uploads", "/settings", "/admin/scan", "/upload"):
            with self.subTest(path=path):
                response = self.client.get(path)
                self.assertEqual(response.status_code, 302)
                self.assertIn("/login", response.location)
                self.assertNotIn(PORTAL, response.location)
                self.assertNotIn(PORTAL, response.get_data(as_text=True))
        self.assertEqual(self.client.get("/submit/wrong-token").status_code, 404)

    def test_public_portal_is_standalone_for_an_authenticated_admin(self):
        self.admin()
        response = self.client.get("/submit/" + PORTAL)
        self.assertEqual(response.status_code, 200)
        self.assertNotIn('id="sidebar"', response.get_data(as_text=True))
        self.assertEqual(response.headers["Referrer-Policy"], "no-referrer")
        self.assertEqual(response.headers["Cache-Control"], "no-store")
        with self.client.session_transaction() as session:
            self.assertIn("_upload_owner", session)
        admin_scan = self.client.get("/admin/scan")
        self.assertEqual(admin_scan.status_code, 200)

    def test_mutating_requests_require_csrf_and_reject_cross_origin(self):
        self.admin()
        for path in ("/api/parse", "/api/uploads/sign", "/api/submit", "/settings", "/logout"):
            with self.subTest(path=path):
                response = self.client.post(path, json={})
                self.assertEqual(response.status_code, 400)
                if path.startswith("/api/"):
                    self.assertIn("error", response.json)
        response = self.client.post(
            "/api/parse", json={"text": "", "portal_token": PORTAL},
            headers={"X-CSRF-Token": CSRF, "Origin": "https://untrusted.example"},
        )
        self.assertEqual(response.status_code, 403)

    def test_null_origin_is_allowed_only_for_same_origin_loopback_development(self):
        headers = {
            "X-CSRF-Token": CSRF,
            "Origin": "null",
            "Sec-Fetch-Site": "same-origin",
        }
        response = self.client.post("/login", data={"password": PASSWORD}, headers=headers)
        self.assertEqual(response.status_code, 302)

        self.authorize_client(self.client)
        response = self.client.post(
            "/login", data={"password": PASSWORD}, headers=headers,
            environ_overrides={"REMOTE_ADDR": "192.0.2.4"},
        )
        self.assertEqual(response.status_code, 403)

        response = self.client.post(
            "/login", data={"password": PASSWORD},
            headers={**headers, "Sec-Fetch-Site": "cross-site"},
        )
        self.assertEqual(response.status_code, 403)

        self.app.config["PRODUCTION"] = True
        response = self.client.post("/login", data={"password": PASSWORD}, headers=headers)
        self.assertEqual(response.status_code, 403)

    def test_admin_login_logout_and_safe_next(self):
        response = self.post("/login?next=//untrusted.example", {"password": PASSWORD})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.location, "/")
        self.assertEqual(self.client.get("/logout").status_code, 405)
        with self.client.session_transaction() as session:
            session["_csrf"] = CSRF
        self.assertEqual(self.post("/logout").status_code, 302)
        self.assertEqual(self.client.get("/trades").status_code, 302)

    def test_local_setup_cannot_be_claimed_remotely_and_requires_ten_characters(self):
        first_app = create_app({**self.config, "ADMIN_PASSWORD_HASH": ""})
        client = first_app.test_client()
        response = client.get("/login", environ_overrides={"REMOTE_ADDR": "192.0.2.4"})
        self.assertEqual(response.status_code, 503)
        response = client.get("/login", headers={"X-Forwarded-For": "192.0.2.4"})
        self.assertEqual(response.status_code, 503)
        self.authorize_client(client)
        response = self.post("/login", {"password": "short", "confirm": "short"}, client=client)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.records("SELECT value FROM settings WHERE key='admin_password'"), [])
        response = self.post("/login", {"password": PASSWORD, "confirm": PASSWORD}, client=client)
        self.assertEqual(response.status_code, 302)
        self.assertEqual(len(self.records("SELECT value FROM settings WHERE key='admin_password'")), 1)

    def test_password_change_invalidates_old_sessions(self):
        self.admin()
        with self.client.session_transaction() as session:
            session["auth_version"] = "old-password-version"
        self.assertEqual(self.client.get("/clients").status_code, 302)

    def test_api_types_lengths_and_body_limit_are_explicit_json_errors(self):
        for value in (None, 42, [], {}):
            response = self.post("/api/parse", json={"portal_token": PORTAL, "text": value})
            self.assertEqual(response.status_code, 400)
            self.assertIn("error", response.json)
        response = self.post("/api/parse", json={"portal_token": PORTAL, "text": "x" * 50001})
        self.assertEqual(response.status_code, 400)
        response = self.post("/api/parse", json={"portal_token": PORTAL, "text": " " * 50001})
        self.assertEqual(response.status_code, 400)
        response = self.post("/api/submit", json=[])
        self.assertEqual(response.status_code, 400)
        response = self.post("/api/parse", data="x", content_type="text/plain")
        self.assertEqual(response.status_code, 415)
        response = self.post("/api/submit", data=b"x" * (4 * 1024 * 1024 + 1), content_type="application/json")
        self.assertEqual(response.status_code, 413)
        self.assertIn("error", response.json)

    def test_api_parse_retains_parser_response_and_requires_a_portal(self):
        response = self.post("/api/parse", json={"text": "INFY 2026-01-02 NSE EQ buy 2 100.00"})
        self.assertEqual(response.status_code, 403)
        response = self.post("/api/parse", json={"portal_token": PORTAL, "text": ""})
        self.assertEqual(response.status_code, 200)
        self.assertIn("trades", response.json)
        self.assertIn("platform", response.json)

    def test_api_preserves_parser_warnings_and_requires_missing_review_fields(self):
        content = "INFY\nBUY\nAvg. 100.00"
        parsed = parse_text(content)
        response = self.post("/api/parse", json={"portal_token": PORTAL, "text": content})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json, parsed)
        self.assertTrue(response.json["warnings"])
        self.assertTrue(all(isinstance(warning, str) for warning in response.json["warnings"]))
        rows = response.json["trades"]
        self.assertTrue(rows)
        self.assertEqual(rows[0]["quantity"], "")
        self.assertEqual(rows[0]["buy_date"], "")
        response = self.submit(self.payload(rows=rows))
        self.assertEqual(response.status_code, 400)
        self.assertIn("quantity", response.json["error"].lower())
        response = self.submit(self.payload(rows=[{**row, "quantity": 1} for row in rows]))
        self.assertEqual(response.status_code, 400)
        self.assertIn("date", response.json["error"].lower())
        self.assertEqual(self.records("SELECT id FROM clients"), [])
        self.assertEqual(self.records("SELECT id FROM trades"), [])
        self.assertEqual(self.records("SELECT id FROM uploads"), [])

    def test_atomic_submission_rejects_all_invalid_numeric_inputs(self):
        for quantity in (None, "", 0, -1, "abc", "NaN", "Infinity", True, [], "0.0000001"):
            with self.subTest(quantity=quantity):
                response = self.submit(self.payload(rows=[dict(ROW), {**ROW, "quantity": quantity}]))
                self.assertEqual(response.status_code, 400, response.get_data(as_text=True))
                self.assertEqual(self.records("SELECT id FROM trades"), [])
                self.assertEqual(self.records("SELECT id FROM clients"), [])
        for price in (-1, "NaN", "Infinity", "bad", False):
            response = self.submit(self.payload(rows=[{**ROW, "buy_price": price}]))
            self.assertEqual(response.status_code, 400)

    def test_dates_are_paired_validated_and_ordered(self):
        for changes in (
            {"buy_date": ""}, {"buy_date": "2026-02-30"}, {"sell_date": ""},
            {"buy_price": "", "buy_date": "2026-01-02"}, {"sell_date": "2026-01-01"},
            {"buy_price": "", "sell_price": "", "buy_date": "", "sell_date": ""},
        ):
            response = self.submit(self.payload(rows=[{**ROW, **changes}]))
            self.assertEqual(response.status_code, 400, response.get_data(as_text=True))
        self.assertEqual(self.records("SELECT id FROM trades"), [])
        response = self.submit(self.payload(rows=[{**ROW, "buy_price": 0}]))
        self.assertEqual(response.status_code, 200)

    def test_public_commission_fields_are_rejected_not_silently_used(self):
        for key in ("commission", "commission_override", "commission_rate", "commission_type"):
            response = self.submit(self.payload(rows=[{**ROW, key: 0}]))
            self.assertEqual(response.status_code, 400)
        self.assertEqual(self.records("SELECT id FROM trades"), [])

    def test_submission_is_idempotent_and_detects_reused_payload_ids(self):
        payload = self.payload()
        first, second = self.submit(payload), self.submit(payload)
        self.assertEqual(first.status_code, 200, first.get_data(as_text=True))
        self.assertEqual(first.json, second.json)
        self.assertEqual(first.json, {"ok": True, "saved": 1, "submission_id": payload["submission_id"]})
        changed = self.submit({**payload, "rows": [{**ROW, "quantity": 5}]})
        self.assertEqual(changed.status_code, 409)
        self.assertEqual(len(self.records("SELECT id FROM trades")), 1)
        self.assertEqual(len(self.records("SELECT id FROM uploads")), 1)

    def test_maximum_batch_is_saved_atomically_without_duplicate_retries(self):
        payload = self.payload(rows=[dict(ROW) for _ in range(500)])
        response = self.submit(payload)
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        self.assertEqual(response.json["saved"], 500)
        self.assertEqual(self.submit(payload).json, response.json)
        self.assertEqual(len(self.records("SELECT id FROM trades")), 500)
        self.assertEqual(self.submit(self.payload(rows=[dict(ROW) for _ in range(501)])).status_code, 400)

    def test_concurrent_retries_create_one_ledger_entry(self):
        with self.client.session_transaction() as session:
            session["_upload_owner"] = "one-browser-owner"
        cookie = self.client.get_cookie("session").value
        payload = self.payload()
        barrier = threading.Barrier(2)

        def gated_client(name):
            barrier.wait(timeout=10)
            return get_or_create_client(name)

        def send_request():
            client = self.app.test_client()
            client.set_cookie("session", cookie)
            return client.post("/api/submit", json=payload, headers={"X-CSRF-Token": CSRF})

        with patch("backend.web.get_or_create_client", side_effect=gated_client):
            with ThreadPoolExecutor(max_workers=2) as executor:
                responses = list(executor.map(lambda _: send_request(), range(2)))
        self.assertEqual([response.status_code for response in responses], [200, 200])
        self.assertEqual(responses[0].json, responses[1].json)
        self.assertEqual(len(self.records("SELECT id FROM trades")), 1)

    def test_case_insensitive_client_matching(self):
        self.assertEqual(self.submit(self.payload(name="Alice Example")).status_code, 200)
        self.assertEqual(self.submit(self.payload(name=" alice   EXAMPLE ")).status_code, 200)
        self.assertEqual(len(self.records("SELECT id FROM clients")), 1)
        self.assertEqual(len(self.records("SELECT id FROM trades")), 2)

    def test_local_direct_upload_and_protected_download(self):
        signed, contents = self.sign()
        self.assertEqual(signed["method"], "PUT")
        self.assertEqual(signed["max_size"], 4 * 1024 * 1024)
        self.assertEqual(self.upload(signed, contents).status_code, 204)
        payload = self.payload(images=[{"path": signed["path"], "receipt": signed["receipt"]}])
        response = self.submit(payload)
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        record = self.records("SELECT client_id,filename FROM uploads")[0]
        self.assertIsNotNone(record["client_id"])
        self.assertEqual(record["filename"], signed["path"])
        self.assertEqual(self.client.get("/uploads/file/" + signed["path"]).status_code, 302)
        self.admin()
        response = self.client.get("/uploads/file/" + signed["path"])
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data, contents)
        response.close()
        self.assertEqual(self.upload(signed, contents).status_code, 409)

    def test_receipts_are_bound_to_the_browser_and_storage_path(self):
        signed, contents = self.sign()
        self.assertEqual(self.upload(signed, contents).status_code, 204)
        other = self.app.test_client()
        self.authorize_client(other)
        response = self.submit(
            self.payload(images=[{"path": signed["path"], "receipt": signed["receipt"]}]), client=other,
        )
        self.assertEqual(response.status_code, 403)
        response = self.submit(self.payload(images=[{"path": "evidence/" + "a" * 32 + ".png", "receipt": signed["receipt"]}]))
        self.assertEqual(response.status_code, 403)
        response = self.submit(self.payload(images=[{"path": signed["path"], "receipt": "forged"}]))
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.records("SELECT id FROM uploads"), [])

    def test_rotation_invalidates_old_link_and_pending_receipts(self):
        signed, contents = self.sign()
        self.admin()
        response = self.post("/settings/rotate-link")
        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.client.get("/submit/" + PORTAL).status_code, 404)
        self.assertEqual(self.submit().status_code, 403)
        self.assertEqual(self.upload(signed, contents).status_code, 403)
        current = self.records("SELECT value FROM settings WHERE key='upload_token'")[0]["value"]
        self.assertNotEqual(current, PORTAL)
        self.assertEqual(self.client.get("/submit/" + current).status_code, 200)

    def test_public_code_is_required_even_when_admin_is_logged_in(self):
        with self.app.app_context():
            set_setting("upload_code", "private-code")
            get_db().commit()
        self.admin()
        self.assertEqual(self.submit().status_code, 403)
        self.assertEqual(self.submit(self.payload(code="private-code")).status_code, 200)
        self.assertEqual(self.submit(self.payload(portal_token="")).status_code, 200)

    def test_incomplete_corrupt_or_too_many_images_do_not_create_trades(self):
        signed, contents = self.sign()
        response = self.submit(self.payload(images=[{"path": signed["path"], "receipt": signed["receipt"]}]))
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.upload(signed, b"x" * len(contents)).status_code, 400)
        self.assertEqual(self.upload(signed, contents[:-1]).status_code, 400)
        response = self.submit(self.payload(images=[{"path": signed["path"], "receipt": signed["receipt"]}] * 11))
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.records("SELECT id FROM trades"), [])
        self.assertEqual(self.records("SELECT id FROM uploads"), [])

    def test_upload_request_rejects_invalid_types_and_sizes(self):
        data = {"portal_token": PORTAL, "filename": "test.png", "content_type": "image/png", "size": 100}
        for changes in ({"size": 0}, {"size": True}, {"size": "100"}, {"size": 4 * 1024 * 1024 + 1},
                        {"content_type": "image/svg+xml"}, {"filename": []}):
            response = self.post("/api/uploads/sign", json={**data, **changes})
            self.assertEqual(response.status_code, 400)

    def test_client_rename_and_merge_keep_submission_relationships(self):
        self.assertEqual(self.submit().status_code, 200)
        cid = self.records("SELECT id FROM clients")[0]["id"]
        self.admin()
        response = self.post(f"/clients/{cid}", {
            "name": "Renamed Client", "commission_type": "", "commission_rate": "",
        })
        self.assertEqual(response.status_code, 302)
        with self.app.app_context():
            target = get_or_create_client("Target Client")
            get_db().commit()
        response = self.client.get(f"/clients/{cid}")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Target Client", response.get_data(as_text=True))
        response = self.post(f"/clients/{cid}/merge", {"target_id": target})
        self.assertEqual(response.status_code, 302)
        upload = self.records("SELECT client_id,client_name FROM uploads")[0]
        self.assertEqual(upload["client_id"], target)
        self.assertEqual(upload["client_name"], "Alice")
        self.assertEqual(self.records("SELECT client_id FROM trades")[0]["client_id"], target)
        self.assertIn("Target Client", self.client.get("/uploads").get_data(as_text=True))

    def test_admin_settings_and_payments_validate_before_writes(self):
        self.admin()
        original = self.records("SELECT value FROM settings WHERE key='default_commission_rate'")[0]["value"]
        for rate in ("-1", "NaN", "nonsense"):
            response = self.post("/settings", {
                "default_commission_type": "flat", "default_commission_rate": rate, "currency": "$",
            })
            self.assertEqual(response.status_code, 400)
        self.assertEqual(self.records("SELECT value FROM settings WHERE key='default_commission_rate'")[0]["value"], original)
        cid = self.submit(self.payload(portal_token="")).json
        self.assertTrue(cid["ok"])
        cid = self.records("SELECT id FROM clients")[0]["id"]
        for changes in ({"amount": "-1"}, {"amount": "0"}, {"amount": "NaN"}, {"paid_on": ""},
                        {"paid_on": "2026-02-30"}, {"amount": "1.234"}):
            response = self.post("/payments", {"client_id": cid, "amount": "10", "paid_on": "2026-01-05", **changes})
            self.assertEqual(response.status_code, 400)
        self.assertEqual(self.records("SELECT id FROM payments"), [])
        response = self.post("/payments", {
            "client_id": cid, "amount": "10", "paid_on": "2026-01-05", "back": "https://untrusted.example",
        })
        self.assertEqual(response.location, "/payments")

    def test_closing_requires_a_price_and_date_and_does_not_silently_succeed(self):
        self.assertEqual(self.submit(self.payload(rows=[{
            **ROW, "sell_price": "", "sell_date": "",
        }])).status_code, 200)
        self.admin()
        tid = self.records("SELECT id FROM trades")[0]["id"]
        for values in ({}, {"sell_price": 0}, {"sell_date": "2026-01-05"},
                       {"sell_price": 10, "sell_date": "2026-01-01"}):
            response = self.post(f"/trades/{tid}/close", values)
            self.assertEqual(response.status_code, 400)
            self.assertIsNone(self.records("SELECT sell_price FROM trades")[0]["sell_price"])
        response = self.post(f"/trades/{tid}/close", {"sell_price": 0, "sell_date": "2026-01-05"})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.records("SELECT sell_price FROM trades")[0]["sell_price"], 0)

    def test_new_commission_rules_are_snapshots_and_legacy_rules_remain_explicit(self):
        self.assertEqual(self.submit().status_code, 200)
        with self.app.app_context():
            set_setting("default_commission_rate", "20")
            get_db().commit()
        self.assertEqual(self.submit().status_code, 200)
        with self.app.app_context():
            trades = query_trades({})
            self.assertEqual([trade["commission"] for trade in reversed(trades)], [Decimal("10"), Decimal("20")])
            get_db().execute(
                "UPDATE trades SET commission_type_snapshot=NULL,commission_rate_snapshot=NULL WHERE id=?",
                (trades[-1]["id"],),
            )
            get_db().commit()
            trades = query_trades({})
            self.assertEqual([trade["commission"] for trade in trades], [Decimal("20"), Decimal("20")])

    def test_export_filters_payments_and_has_correct_all_time_balance_and_safe_strings(self):
        name = "=1+1"
        self.assertEqual(self.submit(self.payload(name=name)).status_code, 200)
        self.assertEqual(self.submit(self.payload(name=name, rows=[{
            **ROW, "stock": "=SUM(1,2)", "sell_price": 200, "notes": "=HYPERLINK(\"https://untrusted.example\")",
            "buy_date": "2026-03-02", "sell_date": "2026-03-05",
        }])).status_code, 200)
        self.assertEqual(self.submit(self.payload(name="Other Private Client")).status_code, 200)
        clients = self.records("SELECT id,name FROM clients ORDER BY id")
        cid, other = clients[0]["id"], clients[1]["id"]
        self.admin()
        for client_id, amount in ((cid, 5), (other, 100)):
            self.assertEqual(self.post("/payments", {
                "client_id": client_id, "amount": amount, "paid_on": "2026-03-10",
            }).status_code, 302)
        response = self.client.get(f"/export/trades.xlsx?client_id={cid}&from=2026-03-01")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.data), data_only=False)
        summary = list(workbook["Summary"].values)
        self.assertEqual(summary[0][-1], "Balance (all time)")
        self.assertEqual(summary[1][-4:], (20, 30, 5, 25))
        self.assertEqual(workbook["Trades"]["B2"].data_type, "s")
        self.assertEqual(workbook["Trades"]["O2"].data_type, "s")
        self.assertEqual(workbook["Summary"]["A2"].data_type, "s")
        self.assertEqual(workbook["Payments"].max_row, 2)
        self.assertEqual(workbook["Payments"]["A2"].value, name)
        workbook.close()

    def test_excel_import_is_atomic_and_rejects_formulas(self):
        self.admin()
        for invalid in ("bad", "=2+2"):
            workbook = Workbook()
            workbook.active.append(["Username", "Stock", "Quantity", "Buy Price", "Buy Date"])
            workbook.active.append(["Good", "INFY", 1, 100, "2026-01-01"])
            workbook.active.append(["Invalid", "INFY", invalid, 100, "2026-01-01"])
            contents = io.BytesIO()
            workbook.save(contents)
            contents.seek(0)
            response = self.post("/import", {"file": (contents, "trades.xlsx")})
            self.assertEqual(response.status_code, 400)
            self.assertEqual(self.records("SELECT id FROM trades"), [])
            self.assertEqual(self.records("SELECT id FROM clients"), [])

    def test_admin_templates_receive_compatible_context(self):
        self.admin()
        for path in ("/", "/trades", "/clients", "/payments", "/uploads", "/settings", "/trades/new", "/admin/scan"):
            with self.subTest(path=path):
                response = self.client.get(path)
                self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        settings = self.client.get("/settings").get_data(as_text=True)
        self.assertIn("/submit/" + PORTAL, settings)

    def test_dashboard_chart_includes_every_positive_client_and_preserves_totals(self):
        for index in range(13):
            response = self.submit(self.payload(name=f"Client {index}", rows=[{**ROW, "quantity": index + 1}]))
            self.assertEqual(response.status_code, 200)
        response = self.submit(self.payload(name="No earnings", rows=[{**ROW, "sell_price": 100}]))
        self.assertEqual(response.status_code, 200)
        self.admin()
        contexts = []

        def capture(_sender, template, context, **_extra):
            contexts.append(context)

        with template_rendered.connected_to(capture, self.app):
            response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        context = contexts[-1]
        chart = context["chart"]
        self.assertEqual(chart["clients"], [f"Client {index}" for index in reversed(range(13))])
        self.assertEqual(chart["client_commission"], [float(5 * quantity) for quantity in range(13, 0, -1)])
        self.assertEqual(sum(chart["client_commission"]), float(context["total"]["commission"]))
        self.assertFalse(context["deployed"])
        self.assertEqual(context["backend_mode"], "local")
        self.assertTrue(context["admin_password_managed"])
        self.assertEqual(context["upload_url"], "http://localhost/submit/" + PORTAL)
        self.assertEqual(context["currency"], "\u20b9")

    def test_large_valid_numbers_do_not_overflow_commission_calculation(self):
        self.admin()
        self.assertEqual(self.post("/settings", {
            "default_commission_type": "turnover_pct", "default_commission_rate": "999999999999", "currency": "$",
        }).status_code, 302)
        response = self.submit(self.payload(rows=[{
            **ROW, "quantity": "999999999999", "buy_price": "999999999999", "sell_price": "999999999999",
        }]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.client.get("/trades").status_code, 200)
        self.assertEqual(self.client.get("/").status_code, 200)

    def test_read_only_migration_preflight_preserves_source_and_links(self):
        signed, contents = self.sign()
        self.assertEqual(self.upload(signed, contents).status_code, 204)
        self.assertEqual(self.submit(self.payload(images=[{
            "path": signed["path"], "receipt": signed["receipt"],
        }])).status_code, 200)
        database_path = Path(self.config["DATABASE"])
        original_database = database_path.read_bytes()
        snapshot = inspect_local(database_path, self.config["UPLOAD_DIR"])
        self.assertEqual(len(snapshot["clients"]), 1)
        self.assertEqual(len(snapshot["assets"]), 1)
        self.assertEqual(snapshot["assets"][0]["upload_id"], snapshot["uploads"][0]["id"])
        self.assertEqual(snapshot["uploads"][0]["client_id"], snapshot["clients"][0]["id"])
        self.assertEqual(database_path.read_bytes(), original_database)
        self.assertEqual((Path(self.config["UPLOAD_DIR"]) / signed["path"]).read_bytes(), contents)
        self.assertNotIn("admin_password", snapshot["settings"])
        self.assertNotIn("upload_token", snapshot["settings"])
        with self.app.app_context(), patch("backend.transfer.get_storage") as storage:
            with self.assertRaises(ValidationError):
                migrate_local(snapshot)
            storage.return_value.upload_bytes.assert_not_called()
        with self.app.app_context():
            get_db().execute("DELETE FROM clients")
            get_db().execute("INSERT INTO clients(name) VALUES ('Alice')")
            get_db().commit()
        snapshot = inspect_local(database_path, self.config["UPLOAD_DIR"])
        self.assertIsNone(snapshot["uploads"][0]["client_id"])

    def test_migration_preflight_rejects_legacy_missing_dates_without_guessing(self):
        self.assertEqual(self.submit().status_code, 200)
        with self.app.app_context():
            get_db().execute("UPDATE trades SET buy_date=NULL")
            get_db().commit()
        with self.assertRaisesRegex(ValidationError, "Trade 1"):
            inspect_local(self.config["DATABASE"], self.config["UPLOAD_DIR"])
        self.assertIsNone(self.records("SELECT buy_date FROM trades")[0]["buy_date"])

    def test_cloud_transfer_rolls_back_upload_failures_but_preserves_ambiguous_commits(self):
        signed, contents = self.sign()
        self.assertEqual(self.upload(signed, contents).status_code, 204)
        self.assertEqual(self.submit(self.payload(images=[{
            "path": signed["path"], "receipt": signed["receipt"],
        }])).status_code, 200)
        snapshot = inspect_local(self.config["DATABASE"], self.config["UPLOAD_DIR"])
        for failure in ("none", "upload", "commit"):
            with self.subTest(failure=failure):
                database, storage = MagicMock(), MagicMock()
                database.postgres = True
                database.execute.return_value.fetchone.return_value = None
                if failure == "upload":
                    storage.upload_bytes.side_effect = StorageUnavailable("Test storage interruption.")
                elif failure == "commit":
                    database.commit.side_effect = psycopg.OperationalError("Test ambiguous commit.")
                with patch("backend.transfer.get_db", return_value=database), patch(
                    "backend.transfer.get_storage", return_value=storage,
                ):
                    if failure == "none":
                        migrate_local(snapshot)
                        database.commit.assert_called_once()
                        storage.remove.assert_not_called()
                    elif failure == "upload":
                        with self.assertRaises(StorageUnavailable):
                            migrate_local(snapshot)
                        database.rollback.assert_called_once()
                        database.commit.assert_not_called()
                        storage.remove.assert_called_once_with([snapshot["assets"][0]["path"]])
                    else:
                        with self.assertRaises(psycopg.OperationalError):
                            migrate_local(snapshot)
                        storage.remove.assert_not_called()

    def test_database_outages_return_safe_errors_without_error_page_recursion(self):
        self.admin()
        with patch("backend.web.get_db", side_effect=psycopg.OperationalError("Do not reveal server secrets")):
            response = self.client.get("/clients")
            self.assertEqual(response.status_code, 503)
            self.assertNotIn("Do not reveal server secrets", response.get_data(as_text=True))
            response = self.submit()
            self.assertEqual(response.status_code, 503)
            self.assertIn("error", response.json)

    def test_cleanup_only_deletes_long_expired_unsubmitted_evidence(self):
        submitted, contents = self.sign()
        self.assertEqual(self.upload(submitted, contents).status_code, 204)
        self.assertEqual(self.submit(self.payload(images=[{
            "path": submitted["path"], "receipt": submitted["receipt"],
        }])).status_code, 200)
        expired, contents = self.sign()
        self.assertEqual(self.upload(expired, contents).status_code, 204)
        pending, contents = self.sign()
        self.assertEqual(self.upload(pending, contents).status_code, 204)
        with self.app.app_context():
            get_db().execute(
                "UPDATE upload_assets SET expires_at=? WHERE path IN (?,?)",
                (int(time.time()) - 2 * 86400, submitted["path"], expired["path"]),
            )
            get_db().commit()
        runner = self.app.test_cli_runner()
        result = runner.invoke(args=["cleanup-uploads", "--dry-run"])
        self.assertEqual(result.exit_code, 0, result.output)
        self.assertTrue((Path(self.config["UPLOAD_DIR"]) / expired["path"]).exists())
        result = runner.invoke(args=["cleanup-uploads"])
        self.assertEqual(result.exit_code, 0, result.output)
        self.assertFalse((Path(self.config["UPLOAD_DIR"]) / expired["path"]).exists())
        self.assertTrue((Path(self.config["UPLOAD_DIR"]) / submitted["path"]).exists())
        self.assertTrue((Path(self.config["UPLOAD_DIR"]) / pending["path"]).exists())
        self.assertEqual(len(self.records("SELECT path FROM upload_assets")), 2)

    def test_request_error_rolls_back_pending_changes(self):
        @self.app.post("/api/test-rollback")
        def rollback_probe():
            get_db().execute("INSERT INTO clients(name) VALUES ('Should Roll Back')")
            raise ValidationError("Deliberate request error.")

        response = self.post("/api/test-rollback")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.records("SELECT id FROM clients"), [])


class LegacyMigrationTests(unittest.TestCase):
    def test_additive_migration_preserves_rows_and_does_not_relink_deleted_clients(self):
        with tempfile.TemporaryDirectory(prefix="commission-legacy-") as directory:
            path = Path(directory) / "legacy.db"
            connection = sqlite3.connect(path)
            connection.executescript("""
                CREATE TABLE settings(key TEXT PRIMARY KEY,value TEXT);
                CREATE TABLE clients(
                    id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT NOT NULL UNIQUE COLLATE NOCASE,
                    phone TEXT,email TEXT,commission_type TEXT,commission_rate REAL,notes TEXT,
                    active INTEGER NOT NULL DEFAULT 1,created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                );
                CREATE TABLE uploads(
                    id INTEGER PRIMARY KEY AUTOINCREMENT,client_name TEXT NOT NULL,filename TEXT,
                    platform TEXT,ocr_text TEXT,trade_count INTEGER DEFAULT 0,created_at TEXT DEFAULT CURRENT_TIMESTAMP
                );
                INSERT INTO clients(id,name) VALUES (7,'Existing');
                INSERT INTO uploads(id,client_name,filename,trade_count) VALUES (12,'existing','old.png',1);
                INSERT INTO settings(key,value) VALUES ('currency','USD');
            """)
            connection.close()
            config = {
                "TESTING": True, "PRODUCTION": False, "BACKEND_MODE": "local",
                "DATABASE": str(path), "SECRET_KEY": "x" * 40, "DATABASE_URL": "",
                "ADMIN_PASSWORD": "", "ADMIN_PASSWORD_HASH": PASSWORD_HASH,
                "PUBLIC_BASE_URL": "", "SESSION_COOKIE_SECURE": False,
            }
            with patch.dict(os.environ, {}, clear=True):
                application = create_app(config)
                with application.app_context():
                    self.assertEqual(get_db().execute("SELECT client_id FROM uploads").fetchone()["client_id"], 7)
                    self.assertEqual(get_db().execute("SELECT value FROM settings WHERE key='currency'").fetchone()["value"], "USD")
                    get_db().execute("DELETE FROM clients WHERE id=7")
                    get_db().execute("INSERT INTO clients(name) VALUES ('Existing')")
                    get_db().commit()
                application = create_app(config)
                with application.app_context():
                    upload = get_db().execute("SELECT * FROM uploads").fetchone()
                    self.assertIsNone(upload["client_id"])
                    self.assertEqual(upload["id"], 12)
                    self.assertEqual(upload["filename"], "old.png")


class CloudContractTests(unittest.TestCase):
    def cloud_config(self):
        return {
            "TESTING": True, "PRODUCTION": True, "BACKEND_MODE": "supabase",
            "SECRET_KEY": "isolated-cloud-secret-" * 3,
            "ADMIN_PASSWORD": "", "ADMIN_PASSWORD_HASH": PASSWORD_HASH,
            "DATABASE_URL": "postgresql://test-user:test-password@pool.example:6543/postgres",
            "SUPABASE_URL": "https://project.example",
            "SUPABASE_SERVICE_ROLE_KEY": "server-only-test-key",
            "SUPABASE_STORAGE_BUCKET": "commission-evidence",
            "PUBLIC_BASE_URL": "https://books.example",
        }

    def test_production_fails_closed_without_secrets_and_never_uses_a_local_secret(self):
        with patch.dict(os.environ, {"VERCEL": "1"}, clear=True):
            with patch("backend.security.os.open") as filesystem_write:
                with self.assertRaisesRegex(RuntimeError, "Production configuration is incomplete"):
                    create_app({"TESTING": True})
                filesystem_write.assert_not_called()
                with patch("backend.database.sqlite3.connect") as sqlite_connect:
                    app = create_app(self.cloud_config())
                    self.assertEqual(app.config["BACKEND_MODE"], "supabase")
                    self.assertTrue(app.config["SESSION_COOKIE_SECURE"])
                    filesystem_write.assert_not_called()
                    sqlite_connect.assert_not_called()

    def test_deployed_flag_requires_vercel_runtime_not_just_cloud_configuration(self):
        cases = (
            ({}, False),
            ({"COMMISSION_ENV": "production"}, False),
            ({"VERCEL": "1", "VERCEL_ENV": "production"}, True),
            ({"VERCEL": "1", "VERCEL_ENV": "preview"}, True),
            ({"VERCEL": "1", "VERCEL_ENV": "development"}, False),
        )
        for environment, deployed in cases:
            with self.subTest(environment=environment), patch.dict(os.environ, environment, clear=True):
                app = create_app(self.cloud_config())
                self.assertIs(app.config["DEPLOYED"], deployed)

    def test_plain_environment_password_has_a_stable_cross_instance_session_version(self):
        with patch.dict(os.environ, {}, clear=True):
            one = Flask("one")
            two = Flask("two")
            config = {**self.cloud_config(), "ADMIN_PASSWORD": PASSWORD, "ADMIN_PASSWORD_HASH": ""}
            configure_app(one, config)
            configure_app(two, config)
            self.assertNotEqual(one.config["ADMIN_PASSWORD_HASH"], two.config["ADMIN_PASSWORD_HASH"])
            self.assertEqual(one.config["ADMIN_AUTH_VERSION"], two.config["ADMIN_AUTH_VERSION"])

    def test_postgres_connections_use_tls_and_disable_prepared_statements(self):
        with patch.dict(os.environ, {}, clear=True):
            app = create_app(self.cloud_config())
        connection = MagicMock()
        connection.closed = False
        with app.app_context(), patch("backend.database.psycopg.connect", return_value=connection) as connect:
            database = get_db()
            self.assertTrue(database.postgres)
            self.assertEqual(connect.call_args.kwargs["sslmode"], "require")
            self.assertIsNone(connect.call_args.kwargs["prepare_threshold"])
            database.execute("SELECT id FROM clients WHERE lower(name)=lower(?)", ("Alice",))
            connection.execute.assert_called_once_with(
                "SELECT id FROM clients WHERE lower(name)=lower(%s)", ("Alice",)
            )
        connection.rollback.assert_called_once()
        connection.close.assert_called_once()

    def test_closed_postgres_connections_do_not_break_error_response_teardown(self):
        connection = MagicMock()
        connection.closed = True
        database = Database(connection, postgres=True)
        database.rollback()
        database.close()
        connection.rollback.assert_not_called()
        connection.close.assert_called_once()

    def test_postgres_result_dates_and_numeric_values_remain_usable(self):
        cursor = MagicMock()
        cursor.fetchone.return_value = {
            "buy_date": date(2026, 1, 2), "created_at": datetime(2026, 1, 2, tzinfo=timezone.utc),
            "buy_price": Decimal("123.456789"),
        }
        row = Result(cursor).fetchone()
        self.assertEqual(row["buy_date"], "2026-01-02")
        self.assertEqual(row["created_at"], "2026-01-02T00:00:00+00:00")
        self.assertEqual(row["buy_price"], Decimal("123.456789"))

    def test_signed_upload_contract_has_no_service_key_and_no_overwrite(self):
        with patch.dict(os.environ, {}, clear=True):
            app = create_app(self.cloud_config())
        path = "evidence/" + "a" * 32 + ".png"
        with app.app_context(), patch("backend.storage.httpx.request") as request:
            request.return_value = httpx.Response(200, json={
                "url": f"/object/upload/sign/commission-evidence/{path}?token=short-lived-capability",
            })
            result = SupabaseStorage().sign_upload(path, "image/png")
            self.assertEqual(result["method"], "PUT")
            self.assertEqual(result["headers"], {"Content-Type": "image/png", "x-upsert": "false"})
            self.assertTrue(result["upload_url"].startswith("https://project.example/storage/v1/object/upload/sign/"))
            self.assertNotIn("server-only-test-key", json.dumps(result))
            self.assertFalse(request.call_args.kwargs["follow_redirects"])
            self.assertEqual(request.call_args.kwargs["headers"]["apikey"], "server-only-test-key")

    def test_cloud_storage_verifies_real_metadata_and_actual_image_bytes(self):
        with patch.dict(os.environ, {}, clear=True):
            app = create_app(self.cloud_config())
        image = png_bytes()
        asset = {"path": "evidence/" + "a" * 32 + ".png", "content_type": "image/png", "byte_size": len(image)}
        with app.app_context(), patch("backend.storage.httpx.request") as request, patch("backend.storage.httpx.stream") as stream:
            request.return_value = httpx.Response(200, json={
                "size": len(image), "content_type": "image/png",
                "metadata": {"size": 99999999, "mimetype": "image/svg+xml"},
            })
            stream.return_value.__enter__.return_value = httpx.Response(200, content=image)
            SupabaseStorage().verify(asset)
            stream.return_value.__enter__.return_value = httpx.Response(200, content=b"x" * len(image))
            with self.assertRaises(ValidationError):
                SupabaseStorage().verify(asset)
            request.return_value = httpx.Response(200, json={
                "size": len(image) + 1, "content_type": "image/png",
                "metadata": {"size": len(image), "mimetype": "image/png"},
            })
            with self.assertRaises(ValidationError):
                SupabaseStorage().verify(asset)

    def test_cloud_bucket_must_really_be_private(self):
        with patch.dict(os.environ, {}, clear=True):
            app = create_app(self.cloud_config())
        with app.app_context(), patch("backend.storage.httpx.request") as request:
            request.return_value = httpx.Response(200, json={
                "public": True, "file_size_limit": 10485760,
                "allowed_mime_types": ["image/png", "image/jpeg", "image/webp"],
            })
            with self.assertRaises(StorageUnavailable):
                SupabaseStorage().check_bucket()
            request.return_value = httpx.Response(503)
            with self.assertRaises(StorageUnavailable):
                SupabaseStorage().check_bucket()

    def test_cloud_status_command_requires_verified_rls_and_storage(self):
        with patch.dict(os.environ, {}, clear=True):
            app = create_app(self.cloud_config())
        database, storage = MagicMock(), MagicMock()
        tables = [
            {"relname": table, "relrowsecurity": True, "browser_access": False}
            for table in ("settings", "clients", "trades", "payments", "uploads", "upload_assets")
        ]
        database.execute.return_value.fetchall.return_value = tables
        with patch("backend.web.get_db", return_value=database), patch(
            "backend.web.get_storage", return_value=storage,
        ):
            result = app.test_cli_runner().invoke(args=["check-cloud"])
            self.assertEqual(result.exit_code, 0, result.output)
            storage.check_bucket.assert_called_once()
            tables[0]["browser_access"] = True
            result = app.test_cli_runner().invoke(args=["check-cloud"])
            self.assertNotEqual(result.exit_code, 0)
            self.assertIn("restrictions are incomplete", result.output)

    def test_number_validation_cannot_accept_nonfinite_or_overprecise_values(self):
        for value in ("NaN", "-Infinity", "Infinity", "1e10000", "1e-100000", "0.0000001", "1,2", "1,,234", "1\u20b92"):
            with self.subTest(value=value), self.assertRaises(ValidationError):
                number(value, "Amount", required=True)


if __name__ == "__main__":
    unittest.main()
